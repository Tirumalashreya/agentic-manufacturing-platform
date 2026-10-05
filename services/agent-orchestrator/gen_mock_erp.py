# One-time generator for mock_erp.xlsx — mirrors real ERP master-data
# structure (separate Customer Master / Product Master / Inventory tables,
# same shape SAP calls KNA1/MARA/MARD). Run once: python gen_mock_erp.py
import random
import openpyxl

random.seed(42)  # same data every time this is regenerated — reproducible tests

wb = openpyxl.Workbook()

# ---------------- Customer Master ----------------
customers = wb.active
customers.title = "Customer Master"
customers.append(["customer_id", "customer_name", "active", "credit_limit"])

customer_names = [
    "ABC Corp", "XYZ Industries", "Old Client Ltd", "Global Tech Solutions",
    "Metro Retail Group", "Sunrise Manufacturing", "Blue Ocean Traders",
    "Nova Electronics", "Pioneer Logistics", "Everest Distributors",
    "Silverline Enterprises", "Redwood Systems", "Falcon Technologies",
    "Harbor Textiles", "Crescent Pharma", "Titan Industrial Supplies",
    "Bluebird Retailers", "Northgate Traders", "Vertex Manufacturing",
    "Coastal Electronics Group", "Ironclad Logistics", "Maple Leaf Distributors",
    "Zenith Enterprises", "Orion Systems", "Cascade Textiles",
    "Summit Pharma Group", "Anchor Industrial Co", "Lighthouse Retail",
    "Frontier Traders", "Quantum Manufacturing",
]
for i, name in enumerate(customer_names, start=1):
    active = not (i % 7 == 0)  # a real spread of inactive accounts, not just 2
    credit_limit = random.choice([50000, 100000, 150000, 200000, 300000, 400000, 500000, 750000, 1000000])
    customers.append([f"C{i:03d}", name, active, credit_limit])

# ---------------- Product Master ----------------
products = wb.create_sheet("Product Master")
products.append(["product_id", "product_name", "price", "configurable", "can_produce", "lead_time_days"])

product_catalog = [
    ("Laptop - Standard", 50000, True, True, 20),
    ("Laptop - Gaming", 95000, True, True, 25),
    ("Laptop - Ultrabook", 78000, True, True, 22),
    ("Desktop - Standard", 35000, False, True, 15),
    ("Desktop - Workstation", 120000, True, True, 30),
    ("Monitor 24in", 8000, False, True, 5),
    ("Monitor 27in", 15000, False, True, 8),
    ("Monitor 32in Curved", 28000, False, True, 10),
    ("Server Rack Unit", 250000, True, True, 45),
    ("Server - Blade", 400000, True, True, 50),
    ("Wireless Mouse", 800, False, True, 3),
    ("Mechanical Keyboard", 3500, False, True, 7),
    ("External SSD 1TB", 6500, False, False, 0),   # vendor-sourced, not produced in-house
    ("External HDD 2TB", 4200, False, False, 0),
    ("Docking Station", 4200, False, True, 10),
    ("Webcam HD", 2200, False, False, 0),
    ("Network Switch 24-port", 18000, False, True, 12),
    ("Router - Enterprise", 22000, True, True, 14),
    ("UPS 1500VA", 9500, False, True, 6),
    ("Printer - Laser", 16000, False, True, 18),
    ("Printer - Multifunction", 24000, True, True, 20),
    ("Tablet - Standard", 28000, True, True, 15),
    ("Projector - HD", 32000, False, True, 16),
    ("Headset - Wireless", 3200, False, True, 5),
    ("Cable Management Kit", 900, False, True, 2),
]
for i, (name, price, configurable, can_produce, lead) in enumerate(product_catalog, start=1):
    products.append([f"P{i:03d}", name, price, configurable, can_produce, lead])

# ---------------- Inventory ----------------
inventory = wb.create_sheet("Inventory")
inventory.append(["product_id", "plant", "stock_qty"])

plants = ["PLANT01", "PLANT02"]
for i in range(1, len(product_catalog) + 1):
    for plant in plants:
        # deliberately include real zero-stock rows, not just high numbers,
        # so shortage scenarios are actually testable
        stock = random.choice([0, 0, 5, 15, 25, 40, 60, 100, 150, 300, 500])
        inventory.append([f"P{i:03d}", plant, stock])

# ---------------- Service Rate Card ----------------
# Test-environment fix (agreed with the user): this generator previously
# only produced Customer/Product/Inventory master data — every sheet below
# existed as a real, already-tested capability in the tools/ layer (real
# service/subscription/procurement/trading workflows, not a stub) but had
# no baseline seed data of its own, so it only ever worked against whatever
# a live session happened to accumulate. Each row below is traced directly
# from a real test's own explicit assertion (see app/tests/test_service.py,
# test_subscription.py, test_procurement.py, test_trading.py,
# test_approvals_sales_order.py) — nothing here is invented beyond what
# those tests already require by name.
service_rates = wb.create_sheet("ServiceRates")
service_rates.append(["service_id", "service_type", "rate_type", "rate", "unit"])
service_rates.append(["SVR001", "Installation", "FIXED", 25000, None])
service_rates.append(["SVR002", "Consulting", "HOURLY", 2500, "hour"])

# ---------------- Subscription Plans ----------------
# Only the two plans an actual find_subscription_plan() lookup exercises
# (test_subscription.py's exact-match and substring-match tests, and
# test_entry_router.py's real end-to-end renewal, which resolves "annual
# software subscription" through this same catalog before acting on the
# customer's real existing subscription row). Basic Support/Cloud Hosting
# are real plan names used elsewhere in the test suite too, but only ever
# passed directly to create_subscription(), never looked up here — adding
# them would be seeding data nothing reads, not a real requirement.
subscription_plans = wb.create_sheet("SubscriptionPlans")
subscription_plans.append(["plan_id", "plan_name", "price", "billing_frequency"])
subscription_plans.append(["PLN002", "Gold Support", 50000, "MONTHLY"])
subscription_plans.append(["PLN004", "Annual Software Subscription", 120000, "ANNUAL"])

# ---------------- BOM (production materials) ----------------
# Laptop - Standard (P001) only — the one product any real test's BOM
# assertion depends on (test_procurement.py's real-shortage-detection
# test). Wireless Mouse (P011) deliberately has NO BOM row: a real test
# (test_check_material_availability_no_bom) asserts has_bom is False for
# it, so every other product is correctly left with no BOM here too,
# matching this catalog's own real "most products aren't manufactured
# in-house from tracked components" shape.
bom = wb.create_sheet("BOM")
bom.append(["bom_id", "product_id", "component_id", "component_name", "qty_per_unit"])
bom.append(["BOM001", "P001", "CHASSIS", "Chassis (generic)", 1])
bom.append(["BOM002", "P001", "CPU", "CPU (generic)", 1])

# ---------------- Component Inventory ----------------
# CHASSIS=25 (30 units needed by the real shortage test -> shortage of
# exactly 5, matching that test's own assertion) and CPU=300 (30 needed,
# comfortably covered -> shortage of exactly 0).
component_inventory = wb.create_sheet("ComponentInventory")
component_inventory.append(["component_id", "component_name", "plant", "stock_qty"])
component_inventory.append(["CHASSIS", "Chassis (generic)", "PLANT01", 25])
component_inventory.append(["CPU", "CPU (generic)", "PLANT01", 300])

# ---------------- Vendors (component/raw-material suppliers) ----------------
# get_vendors_for_component() reads this sheet unconditionally (no
# "if sheet exists" guard, unlike most other lookups here) — it must exist
# in every regenerated workbook. CPU deliberately has exactly one vendor
# (V001); CHASSIS deliberately has two (V004/V005, a real cost-vs-speed
# tradeoff) — both shapes are real, already-tested Procurement Agent
# decisions (test_procurement.py's own vendor-preference tests), not
# arbitrary seed choices.
vendors = wb.create_sheet("Vendors")
vendors.append(["vendor_id", "vendor_name", "component_id", "lead_time_days", "unit_cost"])
vendors.append(["V001", "TechSource Components", "CPU", 10, 8000])
vendors.append(["V004", "ChassisWorks", "CHASSIS", 14, 2500])
vendors.append(["V005", "RapidChassis Express", "CHASSIS", 4, 4200])

# ---------------- Product Vendors (finished-goods trading/drop-ship) ----------------
# Distinct from Vendors above — a vendor for a finished product this
# company sells but doesn't (or can't, right now) manufacture itself.
# Webcam HD (P016) deliberately has two, one of each real route
# (ships_direct_to_customer False = warehouse/Trading, True = drop-ship/
# Third-Party) — test_trading.py's own real assertion. External SSD 1TB
# (P013) has exactly one. Cable Management Kit (P025) has exactly one,
# ships_direct=True — test_approvals_sales_order.py's real prefer_vendor
# test depends on this resolving to a real THIRD_PARTY (not TRADING)
# strategy via graph.py's own strategy_label logic.
product_vendors = wb.create_sheet("ProductVendors")
product_vendors.append(["vendor_id", "vendor_name", "product_id", "lead_time_days", "unit_cost", "ships_direct_to_customer"])
product_vendors.append(["PV001", "Metro Webcam Distributors", "P016", 12, 1800, False])
product_vendors.append(["PV002", "DirectCam Supply Co", "P016", 6, 2000, True])
product_vendors.append(["PV003", "StoragePoint Wholesale", "P013", 8, 5200, False])
product_vendors.append(["PV004", "Nexus Hardware Supply", "P025", 5, 700, True])

wb.save("mock_erp.xlsx")
print(
    f"Wrote mock_erp.xlsx — {len(customer_names)} customers, {len(product_catalog)} products, "
    f"{len(product_catalog) * len(plants)} inventory rows, 2 service rates, 2 subscription plans, "
    f"1 BOM'd product, 2 component-inventory rows, 3 component vendors, 4 product vendors"
)
