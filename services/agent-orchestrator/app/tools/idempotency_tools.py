# app/tools/idempotency_tools.py
import contextlib
import datetime
import fcntl
import json
import openpyxl

MOCK_ERP_PATH = "mock_erp.xlsx"

IDEMPOTENCY_HEADERS = ["idempotency_key", "operation", "result", "created_at"]

# Cross-PROCESS lock (independent-audit P1 fix). get_idempotent_result/
# record_idempotent_result below do an unsynchronized check-then-write
# against mock_erp.xlsx — on their own, two concurrent calls for the same
# idempotency_key can both see "no cached result" and both execute the
# real operation, defeating the whole point of the key. Worse: openpyxl's
# load-modify-save cycle rewrites the ENTIRE workbook file on every save,
# so two concurrent writers to ANY two Excel-backed operations (same key
# or not) can silently clobber each other's appended rows, not just
# double-fire one operation.
#
# fcntl.flock is a real inter-process lock (not just a threading.Lock,
# which only serializes within one process) — stdlib, no new dependency.
# It also correctly serializes multiple THREADS in the same process: each
# caller opens its own file descriptor against the same path, and flock
# blocks on contention regardless of whether the contending holder is a
# different thread or a different process. That dual coverage matters
# here specifically because this service's real, documented deployment
# (see OPERATIONS.md: `uvicorn app.main:app`, no --workers flag) is a
# SINGLE process — so the concurrency this lock actually has to defend
# against today is FastAPI running sync route handlers like create_invoice/
# create_delivery in a thread pool, not separate OS processes. The lock
# still works correctly if that ever changes to multiple worker processes
# on the same host.
#
# Scope (documented, not silently assumed universal): applied to
# create_delivery/create_invoice only — the two functions Fugu's audit
# specifically flagged and the only two _deliver_and_bill calls. The
# other Excel-writing operations in this codebase (create_sales_order,
# create_purchase_order, create_financial_adjustment,
# process_post_sales_action) are NOT wrapped by this fix and can still
# race against each other or against these two now-locked functions.
#
# Hard limit of this whole approach, not just this lock: fcntl.flock is
# per-MACHINE, not distributed — it provides zero protection the moment
# agent-orchestrator ever runs as more than one process across more than
# one host. If that becomes real, the correct fix is migrating this
# mechanism onto the same Postgres-backed IdempotencyKey table every other
# real-time, cross-service write in this codebase already uses
# (inventory-service/order-service/scheduling-service each have their own
# copy of the same claim/replay/409-on-mismatch pattern) — that would also
# remove the underlying single-file-rewrite fragility noted above, which
# this lock only serializes around rather than eliminates. Not built here:
# genuinely out of scope for a P1 concurrency fix on the current
# single-host deployment, but the real path forward if that deployment
# assumption ever stops holding.
_LOCK_PATH = MOCK_ERP_PATH + ".idempotency.lock"


def idempotency_lock():
    """Decorator — wraps the ENTIRE decorated function body (the
    check-cache -> execute real operation -> record-result sequence) in
    one exclusive lock, not just the final file write. A fresh lock
    acquisition happens on every call (contextlib.contextmanager results
    are valid decorators via ContextDecorator, stdlib since 3.2, and
    correctly re-enter a new instance per call — this is not one shared
    lock object reused unsafely across invocations)."""

    @contextlib.contextmanager
    def _cm():
        with open(_LOCK_PATH, "w") as lock_file:
            fcntl.flock(lock_file, fcntl.LOCK_EX)
            try:
                yield
            finally:
                fcntl.flock(lock_file, fcntl.LOCK_UN)

    return _cm()

# The one shared primitive real idempotency actually needs — not a
# framework, just a key/operation -> prior-result lookup. Wired only into
# the operations PLAN.md names (Sales Order, Purchase Order, Delivery,
# Invoice, Credit Memo, Debit Memo): a retried create request for the
# same key returns the exact result it returned the first time instead of
# creating a second real record. Purely additive — idempotency_key is
# optional everywhere it's used, so no existing caller or test is
# affected by not passing one.


def _sheet(wb):
    if "IdempotencyKeys" not in wb.sheetnames:
        sheet = wb.create_sheet("IdempotencyKeys")
        sheet.append(IDEMPOTENCY_HEADERS)
        return sheet
    return wb["IdempotencyKeys"]


def get_idempotent_result(idempotency_key: str | None, operation: str) -> dict | None:
    if not idempotency_key:
        return None
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "IdempotencyKeys" not in wb.sheetnames:
        return None
    sheet = wb["IdempotencyKeys"]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        key, op, result_json, _created_at = row
        if key == idempotency_key and op == operation:
            return json.loads(result_json)
    return None


def record_idempotent_result(idempotency_key: str | None, operation: str, result: dict) -> None:
    if not idempotency_key:
        return
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _sheet(wb)
    sheet.append([idempotency_key, operation, json.dumps(result), datetime.datetime.now().isoformat()])
    wb.save(MOCK_ERP_PATH)
