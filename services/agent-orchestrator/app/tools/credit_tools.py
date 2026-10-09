# app/tools/credit_tools.py
import openpyxl

MOCK_ERP_PATH = "mock_erp.xlsx"

# SalesOrders sheet column indices (see sales_tools.SALES_ORDER_HEADERS)
_SALES_ORDER_ID_COL = 0
_SALES_ORDER_CUSTOMER_ID_COL = 1
_SALES_ORDER_TOTAL_AMOUNT_COL = 7


# Deterministic — no LLM. Existing exposure is the sum of this customer's
# OTHER Sales Orders already on record; a real system would also net out
# what's been paid, but there's no payment tracking yet, so every other
# confirmed order counts as outstanding exposure.
#
# exclude_sales_order_id matters for any caller that creates the Sales
# Order FIRST, then checks credit against it — by the time this runs,
# the order being checked is already a row in the sheet. Without
# excluding it here, its own total_amount got counted once as "existing"
# exposure and a second time via new_order_amount, silently doubling
# every order's apparent exposure and pushing otherwise-valid orders
# onto a credit hold they never should have hit. (The gated path's own
# check_credit, plan_executor.py's _run_check_credit, runs check_credit
# BEFORE create_order and so never needs to pass this at all — it's kept
# for any caller with the opposite ordering.)
def check_credit(
    customer_id: str | None,
    credit_limit: float | None,
    new_order_amount: float,
    exclude_sales_order_id: str | None = None,
) -> dict:
    if customer_id is None or credit_limit is None:
        return {
            "checked": False,
            "reason": "Cannot perform a credit check without a known customer and credit limit.",
        }

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    existing_exposure = 0.0
    if "SalesOrders" in wb.sheetnames:
        sheet = wb["SalesOrders"]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[_SALES_ORDER_ID_COL] == exclude_sales_order_id:
                continue
            if row[_SALES_ORDER_CUSTOMER_ID_COL] == customer_id:
                existing_exposure += row[_SALES_ORDER_TOTAL_AMOUNT_COL] or 0

    total_exposure = existing_exposure + new_order_amount
    within_credit = total_exposure <= credit_limit

    return {
        "checked": True,
        "existing_exposure": existing_exposure,
        "new_order_amount": new_order_amount,
        "total_exposure": total_exposure,
        "credit_limit": credit_limit,
        "within_credit": within_credit,
    }
