# app/tools/post_sales_tools.py
import datetime
import openpyxl

MOCK_ERP_PATH = "mock_erp.xlsx"

POST_SALES_HEADERS = [
    "post_sales_id",
    "action",
    "referenced_order_id",
    "customer_id",
    "customer_name",
    "product_id",
    "product_name",
    "quantity",
    "amount",
    "reason",
    "status",
    "created_at",
    "inspected_at",
    "inspected_by",
]

_ID_PREFIX = {"RETURN": "RET", "CREDIT_MEMO": "CM", "DEBIT_MEMO": "DM"}

# Eligibility policy — a fixed window for now, same pattern as
# pricing_tools' flat discount ceiling before it became a real rule
# table. One value, not per-product/customer, until there's a real need.
RETURN_WINDOW_DAYS = 30


def _post_sales_sheet(wb):
    if "PostSalesActions" not in wb.sheetnames:
        sheet = wb.create_sheet("PostSalesActions")
        sheet.append(POST_SALES_HEADERS)
        return sheet
    return wb["PostSalesActions"]


def _next_post_sales_id(sheet, action: str) -> str:
    prefix = _ID_PREFIX[action]
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith(prefix):
            try:
                max_num = max(max_num, int(str(value)[len(prefix):]))
            except ValueError:
                pass
    return f"{prefix}{max_num + 1:04d}"


def _previously_returned_quantity(referenced_order_id: str, product_id: str, wb=None) -> int:
    wb = wb or openpyxl.load_workbook(MOCK_ERP_PATH)
    if "PostSalesActions" not in wb.sheetnames:
        return 0
    sheet = wb["PostSalesActions"]
    total = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_action, row_order_id, row_product_id, row_quantity, row_status = row[1], row[2], row[5], row[7], row[10]
        # A rejected return never happened — don't hold its quantity
        # against future return attempts on the same order.
        if (
            row_action == "RETURN"
            and row_order_id == referenced_order_id
            and row_product_id == product_id
            and row_status != "REJECTED"
        ):
            total += row_quantity or 0
    return total


def find_sales_order(sales_order_id: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "SalesOrders" not in wb.sheetnames:
        return None
    sheet = wb["SalesOrders"]
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == sales_order_id:
            return dict(zip(headers, row))
    return None


def process_post_sales_action(
    action: str,
    referenced_order_id: str | None,
    customer_id: str | None,
    customer_name: str | None,
    product_id: str | None,
    product_name: str | None,
    quantity: int | None,
    unit_price: float | None,
    reason: str,
) -> dict:
    if product_id is None:
        return {
            "created": False,
            "reason": "Cannot process a post-sales action without a valid product.",
        }
    if quantity is None or quantity <= 0:
        return {
            "created": False,
            "reason": "Cannot process a post-sales action without a valid quantity.",
        }
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)

    # A stated reference must be real — a credit memo/financial adjustment
    # with NO reference at all is still a legitimate case (handled
    # upstream by Process Agent's POST_SALES vs FINANCIAL distinction);
    # this only rejects a reference that was given but doesn't exist.
    if referenced_order_id is not None:
        order = find_sales_order(referenced_order_id)
        if order is None:
            return {
                "created": False,
                "reason": f"Referenced order {referenced_order_id} does not exist.",
            }
        # Quantity validation is only meaningful against the right product
        # on that order — checking "quantity <= ordered" against the wrong
        # product would be nonsense, so this must be checked first.
        if order["product_id"] != product_id:
            return {
                "created": False,
                "reason": (
                    f"Product {product_id} was not part of order "
                    f"{referenced_order_id} (that order was for {order['product_id']})."
                ),
            }
        if quantity > order["quantity"]:
            return {
                "created": False,
                "reason": (
                    f"Requested quantity {quantity} exceeds the "
                    f"{order['quantity']} units on order {referenced_order_id}."
                ),
            }
        # Cumulative check — each individual return can look fine against
        # the order total on its own, but a second return still can't push
        # the running total past what was actually sold.
        if action == "RETURN":
            already_returned = _previously_returned_quantity(referenced_order_id, product_id, wb=wb)
            if already_returned + quantity > order["quantity"]:
                return {
                    "created": False,
                    "reason": (
                        f"Requested return of {quantity} plus {already_returned} "
                        f"already returned would exceed the {order['quantity']} "
                        f"units on order {referenced_order_id}."
                    ),
                }
            # Eligibility window — an order from a year ago isn't
            # returnable just because the quantity/product line up.
            order_age_days = (
                datetime.date.today()
                - datetime.date.fromisoformat(order["created_at"][:10])
            ).days
            if order_age_days > RETURN_WINDOW_DAYS:
                return {
                    "created": False,
                    "reason": (
                        f"Order {referenced_order_id} was placed {order_age_days} days ago, "
                        f"beyond the {RETURN_WINDOW_DAYS}-day return window."
                    ),
                }

    amount = round(unit_price * quantity, 2) if unit_price is not None else None

    sheet = _post_sales_sheet(wb)

    post_sales_id = _next_post_sales_id(sheet, action)
    created_at = datetime.datetime.now().isoformat()
    # A return isn't final the moment it's requested — the goods have to
    # actually come back and be inspected before stock/credit is granted.
    # Credit/debit memos are financial-only, no physical goods involved,
    # so they stay immediate.
    status = "PENDING_INSPECTION" if action == "RETURN" else "ISSUED"

    sheet.append([
        post_sales_id,
        action,
        referenced_order_id,
        customer_id,
        customer_name,
        product_id,
        product_name,
        quantity,
        amount,
        reason,
        status,
        created_at,
        None,
        None,
    ])
    wb.save(MOCK_ERP_PATH)

    return {
        "created": True,
        "post_sales_id": post_sales_id,
        "amount": amount,
        "status": status,
    }


def get_post_sales_action(post_sales_id: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "PostSalesActions" not in wb.sheetnames:
        return None
    sheet = wb["PostSalesActions"]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == post_sales_id:
            return dict(zip(POST_SALES_HEADERS, row))
    return None


# The inspection step. Only a PENDING_INSPECTION return can be resolved;
# only APPROVED actually moves stock — a rejected return changes nothing
# physically, matching real RMA handling (goods that fail inspection don't
# go back into sellable inventory here).
def resolve_return_inspection(post_sales_id: str, decision: str, inspected_by: str = "SYSTEM") -> dict:
    if decision not in ("APPROVED", "REJECTED"):
        return {"resolved": False, "reason": f"Invalid decision: {decision}"}

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "PostSalesActions" not in wb.sheetnames:
        return {"resolved": False, "reason": f"No post-sales action found with id {post_sales_id}"}
    sheet = wb["PostSalesActions"]

    action_col = POST_SALES_HEADERS.index("action") + 1
    product_id_col = POST_SALES_HEADERS.index("product_id") + 1
    quantity_col = POST_SALES_HEADERS.index("quantity") + 1
    status_col = POST_SALES_HEADERS.index("status") + 1
    inspected_at_col = POST_SALES_HEADERS.index("inspected_at") + 1
    inspected_by_col = POST_SALES_HEADERS.index("inspected_by") + 1

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value != post_sales_id:
            continue
        if row[action_col - 1].value != "RETURN":
            return {"resolved": False, "reason": f"{post_sales_id} is not a RETURN action."}
        current_status = row[status_col - 1].value
        if current_status != "PENDING_INSPECTION":
            return {
                "resolved": False,
                "reason": f"{post_sales_id} is not pending inspection (status: {current_status}).",
            }

        product_id = row[product_id_col - 1].value
        quantity = row[quantity_col - 1].value

        sheet.cell(row=row_idx, column=status_col, value=decision)
        sheet.cell(row=row_idx, column=inspected_at_col, value=datetime.datetime.now().isoformat())
        sheet.cell(row=row_idx, column=inspected_by_col, value=inspected_by)

        if decision == "APPROVED":
            inventory = wb["Inventory"]
            for inv_row in inventory.iter_rows(min_row=2):
                pid_cell, _plant_cell, qty_cell = inv_row
                if pid_cell.value == product_id:
                    qty_cell.value = (qty_cell.value or 0) + quantity
                    break

        wb.save(MOCK_ERP_PATH)
        return {"resolved": True, "post_sales_id": post_sales_id, "status": decision}

    return {"resolved": False, "reason": f"No post-sales action found with id {post_sales_id}"}
