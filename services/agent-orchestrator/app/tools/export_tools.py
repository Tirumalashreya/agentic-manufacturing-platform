# app/tools/export_tools.py
import datetime
import openpyxl

from app.erp.mock_erp_adapter import default_erp_adapter as erp
from app.tools.audit_tools import record_audit_event

MOCK_ERP_PATH = "mock_erp.xlsx"

# ExportShipment/ExportDocument/CustomsDeclaration/FreightBooking are all
# 1:1 with a single ExportOrder — folded directly into its row rather than
# kept as separate near-empty sheets, same pattern as TradingOrders (every
# lifecycle field in-row) and ServiceOrders/Projects before it.
EXPORT_ORDER_HEADERS = [
    "export_order_id",
    "customer_id",
    "customer_name",
    "product_id",
    "product_name",
    "quantity",
    "unit_price",
    "total_value",
    "destination_country",
    "incoterm",
    "proforma_invoice_id",
    "payment_reference",
    "fulfillment_source",
    "production_order_id",
    "reservation_id",
    "customs_declaration_no",
    "freight_booking_ref",
    "bl_awb_number",
    "delivery_id",
    "export_invoice_id",
    "status",
    "created_at",
    "shipped_at",
    "closed_at",
]

EXPORT_PROFORMA_TERMS_DAYS = 0  # informational, due on issue — not a payment demand
EXPORT_INVOICE_PAYMENT_TERMS_DAYS = 30  # same flat Net 30 used elsewhere (billing_agent, service_tools)

# Real, person/operationally-triggered transitions, same idiom as
# TradingOrders/ServiceOrders/ProjectMilestones. ALLOCATED and
# IN_PRODUCTION are alternative predecessors of DOCUMENTED (stock vs.
# production fulfillment), decided in graph.py's
# allocate_or_produce_export_order() since that step needs the same
# procurement-agent reuse MTO/CTO/Project already use.
_TRANSITIONS = {
    "CONFIRMED": {"PAYMENT_VERIFIED"},
    "PAYMENT_VERIFIED": {"ALLOCATED", "IN_PRODUCTION"},
    "ALLOCATED": {"DOCUMENTED"},
    "IN_PRODUCTION": {"DOCUMENTED"},
    "DOCUMENTED": {"FREIGHT_BOOKED"},
    "FREIGHT_BOOKED": {"SHIPPED"},
    "SHIPPED": {"INVOICED"},
    "INVOICED": {"CLOSED"},
}


def _export_orders_sheet(wb):
    if "ExportOrders" not in wb.sheetnames:
        sheet = wb.create_sheet("ExportOrders")
        sheet.append(EXPORT_ORDER_HEADERS)
        return sheet
    return wb["ExportOrders"]


def _next_export_order_id(sheet) -> str:
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith("EXP"):
            try:
                max_num = max(max_num, int(str(value)[3:]))
            except ValueError:
                pass
    return f"EXP{max_num + 1:04d}"


def find_export_order(export_order_id: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ExportOrders" not in wb.sheetnames:
        return None
    sheet = wb["ExportOrders"]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == export_order_id:
            return dict(zip(EXPORT_ORDER_HEADERS, row))
    return None


# Full history for the Export Sales channel view — includes CLOSED orders,
# unlike list_open_export_orders() which powers the Active Workflows queue.
def list_export_orders() -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ExportOrders" not in wb.sheetnames:
        return []
    sheet = wb["ExportOrders"]
    rows = [dict(zip(EXPORT_ORDER_HEADERS, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    rows.reverse()
    return rows


def list_open_export_orders() -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ExportOrders" not in wb.sheetnames:
        return []
    sheet = wb["ExportOrders"]
    return [
        dict(zip(EXPORT_ORDER_HEADERS, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[EXPORT_ORDER_HEADERS.index("status")] != "CLOSED"
    ]


def _set_fields(export_order_id: str, **fields) -> bool:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ExportOrders" not in wb.sheetnames:
        return False
    sheet = wb["ExportOrders"]
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value != export_order_id:
            continue
        for name, value in fields.items():
            sheet.cell(row=row_idx, column=EXPORT_ORDER_HEADERS.index(name) + 1, value=value)
        wb.save(MOCK_ERP_PATH)
        return True
    return False


def _advance(export_order_id: str, new_status: str, error_label: str, **fields) -> dict:
    order = find_export_order(export_order_id)
    if order is None:
        return {"updated": False, "reason": f"No export order found with id {export_order_id}"}
    allowed = _TRANSITIONS.get(order["status"], set())
    if new_status not in allowed:
        return {
            "updated": False,
            "reason": (
                f"Cannot {error_label} {export_order_id} from status {order['status']}. "
                f"Allowed next status(es): {sorted(allowed) or 'none — terminal'}."
            ),
        }
    _set_fields(export_order_id, status=new_status, **fields)
    record_audit_event("ExportOrder", export_order_id, f"EXPORT_ORDER_{new_status}", fields)
    return {"updated": True, "export_order_id": export_order_id, "status": new_status}


# Creates the order and its proforma invoice in one shot — reuses
# create_invoice() exactly as the roadmap specified (no dedicated
# ExportBillingAgent). destination_country is required (rejected if
# missing, never invented by the agent); incoterm is optional.
def create_export_order(
    customer_id: str | None,
    customer_name: str | None,
    product_id: str | None,
    product_name: str | None,
    quantity: int | None,
    unit_price: float | None,
    destination_country: str | None,
    incoterm: str | None,
) -> dict:
    if not destination_country:
        return {"created": False, "reason": "Cannot create an export order without a destination country."}
    if product_id is None or unit_price is None:
        return {"created": False, "reason": "Cannot create an export order without a valid product and price."}
    if quantity is None or quantity <= 0:
        return {"created": False, "reason": "Cannot create an export order without a valid quantity."}

    total_value = round(unit_price * quantity, 2)

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _export_orders_sheet(wb)
    export_order_id = _next_export_order_id(sheet)
    created_at = datetime.datetime.now().isoformat()

    sheet.append([
        export_order_id, customer_id, customer_name, product_id, product_name,
        quantity, unit_price, total_value, destination_country, incoterm,
        None, None, None, None, None, None, None, None, None, None,
        "CONFIRMED", created_at, None, None,
    ])
    wb.save(MOCK_ERP_PATH)

    invoice = erp.create_invoice(
        sales_order_id=None,
        delivery_id=None,
        customer_id=customer_id,
        customer_name=customer_name,
        quantity_billed=quantity,
        unit_price=unit_price,
        payment_terms_days=EXPORT_PROFORMA_TERMS_DAYS,
        notes=f"Proforma invoice for export order {export_order_id} to {destination_country}.",
    )
    proforma_invoice_id = invoice["invoice_id"] if invoice["created"] else None
    _set_fields(export_order_id, proforma_invoice_id=proforma_invoice_id)
    record_audit_event("ExportOrder", export_order_id, "EXPORT_ORDER_CREATED", {
        "destination_country": destination_country, "total_value": total_value,
    })

    return {
        "created": True,
        "export_order_id": export_order_id,
        "total_value": total_value,
        "proforma_invoice_id": proforma_invoice_id,
        "status": "CONFIRMED",
    }


# Records the stock-vs-production fulfillment decision made in graph.py's
# allocate_or_produce_export_order() (that decision needs procurement_agent
# reuse, so it lives alongside dispatch_trading_order() there, not here) —
# this just persists the outcome and advances the status, same
# create/decide-elsewhere-record-here split as attach_production_order().
def set_export_fulfillment(
    export_order_id: str,
    fulfillment_source: str,
    new_status: str,
    reservation_id: str | None = None,
    production_order_id: str | None = None,
) -> dict:
    return _advance(
        export_order_id, new_status, "allocate",
        fulfillment_source=fulfillment_source,
        reservation_id=reservation_id,
        production_order_id=production_order_id,
    )


# LC / advance-payment verification gate. A real system would check this
# against an actual bank/LC record — there isn't one here, so this
# requires an explicit, non-empty reference rather than faking automatic
# approval, same "never fake a gate" principle as resolve_return_inspection.
def verify_export_payment(export_order_id: str, payment_reference: str) -> dict:
    if not payment_reference:
        return {"updated": False, "reason": "A real payment/LC reference is required to verify export payment."}
    return _advance(export_order_id, "PAYMENT_VERIFIED", "verify payment for", payment_reference=payment_reference)


def record_documentation(export_order_id: str, customs_declaration_no: str) -> dict:
    if not customs_declaration_no:
        return {"updated": False, "reason": "A real customs declaration number is required."}
    return _advance(export_order_id, "DOCUMENTED", "record documentation for", customs_declaration_no=customs_declaration_no)


def record_freight_booking(export_order_id: str, freight_booking_ref: str) -> dict:
    if not freight_booking_ref:
        return {"updated": False, "reason": "A real freight booking reference is required."}
    return _advance(export_order_id, "FREIGHT_BOOKED", "record freight booking for", freight_booking_ref=freight_booking_ref)


# Creates a real Delivery — reuses create_delivery() exactly as the
# roadmap specified.
def ship_export_order(export_order_id: str, bl_awb_number: str) -> dict:
    if not bl_awb_number:
        return {"updated": False, "reason": "A real BL/AWB number is required to ship."}

    order = find_export_order(export_order_id)
    if order is None:
        return {"updated": False, "reason": f"No export order found with id {export_order_id}"}
    if order["status"] != "FREIGHT_BOOKED":
        return {
            "updated": False,
            "reason": f"Cannot ship {export_order_id} from status {order['status']} (expected FREIGHT_BOOKED).",
        }

    delivery = erp.create_delivery(
        sales_order_id=None,
        customer_id=order["customer_id"],
        customer_name=order["customer_name"],
        product_id=order["product_id"],
        product_name=order["product_name"],
        quantity=order["quantity"],
    )
    if not delivery["created"]:
        return {"updated": False, "reason": f"Shipment failed: {delivery['reason']}"}

    result = _advance(
        export_order_id, "SHIPPED", "ship",
        bl_awb_number=bl_awb_number,
        delivery_id=delivery["delivery_id"],
        shipped_at=datetime.datetime.now().isoformat(),
    )
    result["delivery_id"] = delivery["delivery_id"]
    return result


# The commercial/export invoice — distinct from the proforma issued at
# creation, same real-world distinction (a proforma is not a demand for
# payment; the commercial invoice issued after shipment is). Reuses
# create_invoice() again, no dedicated agent.
def invoice_export_order(export_order_id: str) -> dict:
    order = find_export_order(export_order_id)
    if order is None:
        return {"updated": False, "reason": f"No export order found with id {export_order_id}"}
    if order["status"] != "SHIPPED":
        return {
            "updated": False,
            "reason": f"Cannot invoice {export_order_id} from status {order['status']} (expected SHIPPED).",
        }

    invoice = erp.create_invoice(
        sales_order_id=None,
        delivery_id=order["delivery_id"],
        customer_id=order["customer_id"],
        customer_name=order["customer_name"],
        quantity_billed=order["quantity"],
        unit_price=order["unit_price"],
        payment_terms_days=EXPORT_INVOICE_PAYMENT_TERMS_DAYS,
        notes=f"Export commercial invoice for {export_order_id} ({order['incoterm']} to {order['destination_country']}).",
    )
    if not invoice["created"]:
        return {"updated": False, "reason": f"Invoicing failed: {invoice['reason']}"}

    result = _advance(export_order_id, "INVOICED", "invoice", export_invoice_id=invoice["invoice_id"])
    result["export_invoice_id"] = invoice["invoice_id"]
    return result


# Realizes payment on the commercial invoice (the first real use of
# mark_invoice_paid()) and closes the order. Terminal.
def close_export_order(export_order_id: str) -> dict:
    order = find_export_order(export_order_id)
    if order is None:
        return {"closed": False, "reason": f"No export order found with id {export_order_id}"}
    if order["status"] != "INVOICED":
        return {
            "closed": False,
            "reason": f"Cannot close {export_order_id} from status {order['status']} (expected INVOICED).",
        }

    payment = erp.mark_invoice_paid(order["export_invoice_id"])
    if not payment["updated"]:
        return {"closed": False, "reason": f"Payment realisation failed: {payment['reason']}"}

    result = _advance(export_order_id, "CLOSED", "close", closed_at=datetime.datetime.now().isoformat())
    return {"closed": result["updated"], "export_order_id": export_order_id, "status": result.get("status")}
