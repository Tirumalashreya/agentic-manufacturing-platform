# app/tools/service_tools.py
import datetime
import openpyxl

from app.erp.mock_erp_adapter import default_erp_adapter as erp
from app.tools.audit_tools import record_audit_event

MOCK_ERP_PATH = "mock_erp.xlsx"

SERVICE_ORDER_HEADERS = [
    "service_order_id",
    "customer_id",
    "customer_name",
    "service_type",
    "scope_description",
    "rate_type",
    "rate",
    "estimated_quantity",
    "estimated_amount",
    "status",
    "created_at",
    "actual_hours",
    "billable_amount",
    "signed_off_by",
    "completed_at",
    "invoice_id",
]

# Flat Net 30 — same value billing_agent always lands on for product
# deliveries (this system has no customer-specific payment term yet), and
# a service completion has no partial-delivery nuance for an LLM to weigh
# in on, so this skips the agent call entirely rather than asking a model
# to always answer "30".
SERVICE_PAYMENT_TERMS_DAYS = 30

# Real transitions a person can explicitly trigger — no time-based
# auto-progression, same principle as Production Orders never completing
# on their own. COMPLETED and CANCELLED are terminal.
_VALID_STATUS_TRANSITIONS = {
    "REQUESTED": {"SCHEDULED", "CANCELLED"},
    "SCHEDULED": {"IN_PROGRESS", "CANCELLED"},
    "IN_PROGRESS": {"COMPLETED"},
}


def _service_orders_sheet(wb):
    if "ServiceOrders" not in wb.sheetnames:
        sheet = wb.create_sheet("ServiceOrders")
        sheet.append(SERVICE_ORDER_HEADERS)
        return sheet
    return wb["ServiceOrders"]


def _next_service_order_id(sheet) -> str:
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith("SVC"):
            try:
                max_num = max(max_num, int(str(value)[3:]))
            except ValueError:
                pass
    return f"SVC{max_num + 1:04d}"


# Deterministic rate lookup against the real ServiceRates master — same
# exact-match-then-substring philosophy as find_subscription_plan(). No
# rate is invented for a service_type that isn't on file.
def find_service_rate(service_type: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ServiceRates" not in wb.sheetnames:
        return None
    sheet = wb["ServiceRates"]
    needle = service_type.strip().lower()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        service_id, row_service_type, rate_type, rate, unit = row
        if row_service_type.strip().lower() == needle:
            return {"service_id": service_id, "service_type": row_service_type, "rate_type": rate_type, "rate": rate, "unit": unit}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        service_id, row_service_type, rate_type, rate, unit = row
        haystack = row_service_type.strip().lower()
        if needle in haystack or haystack in needle:
            return {"service_id": service_id, "service_type": row_service_type, "rate_type": rate_type, "rate": rate, "unit": unit}

    return None


def find_service_order(service_order_id: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ServiceOrders" not in wb.sheetnames:
        return None
    sheet = wb["ServiceOrders"]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == service_order_id:
            return dict(zip(SERVICE_ORDER_HEADERS, row))
    return None


# Full history for the Services Sales channel view — includes COMPLETED/
# CANCELLED, unlike list_open_service_orders() which powers Active Workflows.
def list_service_orders() -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ServiceOrders" not in wb.sheetnames:
        return []
    sheet = wb["ServiceOrders"]
    rows = [dict(zip(SERVICE_ORDER_HEADERS, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    rows.reverse()
    return rows


def list_open_service_orders() -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ServiceOrders" not in wb.sheetnames:
        return []
    sheet = wb["ServiceOrders"]
    status_idx = SERVICE_ORDER_HEADERS.index("status")
    return [
        dict(zip(SERVICE_ORDER_HEADERS, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[status_idx] not in ("COMPLETED", "CANCELLED")
    ]


# Records a real, static REQUESTED order with a real estimated amount when
# a rate is on file. No milestone/execution progression is simulated —
# that's future work (see update_service_order_status for the real,
# explicitly-triggered status transitions that do exist).
#
# estimated_hours only matters for an HOURLY-rated service_type; if the
# customer didn't state an estimate, estimated_amount stays None rather
# than guessing a number of hours.
def create_service_order(
    customer_id: str | None,
    customer_name: str | None,
    service_type: str,
    scope_description: str,
    estimated_hours: int | None = None,
) -> dict:
    if not scope_description:
        return {
            "created": False,
            "reason": "Cannot create a service order without a scope description.",
        }

    rate_info = find_service_rate(service_type)
    if rate_info is None:
        rate_type = rate = estimated_quantity = estimated_amount = None
    elif rate_info["rate_type"] == "FIXED":
        rate_type = "FIXED"
        rate = rate_info["rate"]
        estimated_quantity = 1
        estimated_amount = rate
    else:
        rate_type = "HOURLY"
        rate = rate_info["rate"]
        estimated_quantity = estimated_hours
        estimated_amount = round(rate * estimated_hours, 2) if estimated_hours else None

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _service_orders_sheet(wb)

    service_order_id = _next_service_order_id(sheet)
    created_at = datetime.datetime.now().isoformat()
    status = "REQUESTED"

    sheet.append([
        service_order_id,
        customer_id,
        customer_name,
        service_type,
        scope_description,
        rate_type,
        rate,
        estimated_quantity,
        estimated_amount,
        status,
        created_at,
        None,
        None,
        None,
        None,
        None,
    ])
    wb.save(MOCK_ERP_PATH)
    record_audit_event("ServiceOrder", service_order_id, "SERVICE_ORDER_CREATED", {"service_type": service_type})

    return {
        "created": True,
        "service_order_id": service_order_id,
        "rate_type": rate_type,
        "rate": rate,
        "estimated_amount": estimated_amount,
        "status": status,
    }


# The explicit, person-triggered progression step. Only a real transition
# in _VALID_STATUS_TRANSITIONS is allowed — no skipping stages, no
# resurrecting a terminal order.
def update_service_order_status(service_order_id: str, new_status: str) -> dict:
    if new_status not in {"SCHEDULED", "IN_PROGRESS", "COMPLETED", "CANCELLED"}:
        return {"updated": False, "reason": f"Invalid status: {new_status}"}

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ServiceOrders" not in wb.sheetnames:
        return {"updated": False, "reason": f"No service order found with id {service_order_id}"}
    sheet = wb["ServiceOrders"]

    status_col = SERVICE_ORDER_HEADERS.index("status") + 1

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value != service_order_id:
            continue
        current_status = row[status_col - 1].value
        allowed = _VALID_STATUS_TRANSITIONS.get(current_status, set())
        if new_status not in allowed:
            return {
                "updated": False,
                "reason": (
                    f"Cannot move {service_order_id} from {current_status} to {new_status}. "
                    f"Allowed next status(es): {sorted(allowed) or 'none — terminal'}."
                ),
            }

        sheet.cell(row=row_idx, column=status_col, value=new_status)
        wb.save(MOCK_ERP_PATH)
        record_audit_event("ServiceOrder", service_order_id, f"SERVICE_ORDER_{new_status}")
        return {"updated": True, "service_order_id": service_order_id, "status": new_status}

    return {"updated": False, "reason": f"No service order found with id {service_order_id}"}


# The completion/billing step. Only a real IN_PROGRESS engagement can be
# completed — same terminal-transition guard as update_service_order_status,
# reused here rather than duplicated. An HOURLY order needs real
# actual_hours to bill (the original estimated_hours was only ever a
# guess); a FIXED order bills the flat rate regardless of hours worked.
# No rate on file at all still completes the engagement (the work still
# happened) but leaves billing for manual handling, same honest fallback
# create_service_order already uses when a rate can't be found.
def complete_service_order(
    service_order_id: str,
    actual_hours: int | None = None,
    signed_off_by: str = "SYSTEM",
) -> dict:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ServiceOrders" not in wb.sheetnames:
        return {"completed": False, "reason": f"No service order found with id {service_order_id}"}
    sheet = wb["ServiceOrders"]

    col = {name: SERVICE_ORDER_HEADERS.index(name) + 1 for name in SERVICE_ORDER_HEADERS}

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value != service_order_id:
            continue

        current_status = row[col["status"] - 1].value
        allowed = _VALID_STATUS_TRANSITIONS.get(current_status, set())
        if "COMPLETED" not in allowed:
            return {
                "completed": False,
                "reason": (
                    f"Cannot complete {service_order_id} from status {current_status}. "
                    f"Allowed next status(es): {sorted(allowed) or 'none — terminal'}."
                ),
            }

        rate_type = row[col["rate_type"] - 1].value
        rate = row[col["rate"] - 1].value
        customer_id = row[col["customer_id"] - 1].value
        customer_name = row[col["customer_name"] - 1].value
        service_type = row[col["service_type"] - 1].value

        if rate_type == "HOURLY":
            if not actual_hours or actual_hours <= 0:
                return {
                    "completed": False,
                    "reason": f"{service_order_id} is billed HOURLY — actual_hours is required to complete it.",
                }
            billable_amount = round(rate * actual_hours, 2)
            invoice_quantity = actual_hours
        elif rate_type == "FIXED":
            billable_amount = rate
            invoice_quantity = 1
        else:
            billable_amount = None
            invoice_quantity = None

        completed_at = datetime.datetime.now().isoformat()
        invoice_id = None
        if billable_amount is not None:
            invoice = erp.create_invoice(
                sales_order_id=None,
                delivery_id=None,
                customer_id=customer_id,
                customer_name=customer_name,
                quantity_billed=invoice_quantity,
                unit_price=rate,
                payment_terms_days=SERVICE_PAYMENT_TERMS_DAYS,
                notes=f"Service {service_order_id} ({service_type}) completed, signed off by {signed_off_by}.",
            )
            if invoice["created"]:
                invoice_id = invoice["invoice_id"]

        sheet.cell(row=row_idx, column=col["status"], value="COMPLETED")
        sheet.cell(row=row_idx, column=col["actual_hours"], value=actual_hours)
        sheet.cell(row=row_idx, column=col["billable_amount"], value=billable_amount)
        sheet.cell(row=row_idx, column=col["signed_off_by"], value=signed_off_by)
        sheet.cell(row=row_idx, column=col["completed_at"], value=completed_at)
        sheet.cell(row=row_idx, column=col["invoice_id"], value=invoice_id)
        wb.save(MOCK_ERP_PATH)
        record_audit_event("ServiceOrder", service_order_id, "SERVICE_ORDER_COMPLETED", {
            "billable_amount": billable_amount, "invoice_id": invoice_id,
        })

        return {
            "completed": True,
            "service_order_id": service_order_id,
            "billable_amount": billable_amount,
            "invoice_id": invoice_id,
            "status": "COMPLETED",
        }

    return {"completed": False, "reason": f"No service order found with id {service_order_id}"}
