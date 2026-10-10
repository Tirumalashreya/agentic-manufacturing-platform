# app/tools/billing_tools.py
import datetime
import openpyxl

from app.tools.idempotency_tools import get_idempotent_result, idempotency_lock, record_idempotent_result

MOCK_ERP_PATH = "mock_erp.xlsx"

INVOICE_HEADERS = [
    "invoice_id",
    "sales_order_id",
    "delivery_id",
    "customer_id",
    "customer_name",
    "quantity_billed",
    "unit_price",
    "invoice_amount",
    "payment_terms_days",
    "due_date",
    "status",
    "notes",
    "created_at",
]


def _invoices_sheet(wb):
    if "Invoices" not in wb.sheetnames:
        sheet = wb.create_sheet("Invoices")
        sheet.append(INVOICE_HEADERS)
        return sheet
    return wb["Invoices"]


def _next_invoice_id(sheet) -> str:
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith("INV"):
            try:
                max_num = max(max_num, int(str(value)[3:]))
            except ValueError:
                pass
    return f"INV{max_num + 1:04d}"


@idempotency_lock()
def create_invoice(
    sales_order_id: str,
    delivery_id: str,
    customer_id: str | None,
    customer_name: str | None,
    quantity_billed: int,
    unit_price: float,
    payment_terms_days: int,
    notes: str,
    idempotency_key: str | None = None,
) -> dict:
    cached = get_idempotent_result(idempotency_key, "create_invoice")
    if cached is not None:
        return cached

    if quantity_billed is None or quantity_billed <= 0 or unit_price is None:
        return {
            "created": False,
            "reason": "Cannot create an invoice without a valid billed quantity and unit price.",
        }

    # Computed here, not accepted as a caller-supplied value — same
    # principle as create_quotation()/create_sales_order().
    invoice_amount = round(unit_price * quantity_billed, 2)

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _invoices_sheet(wb)

    invoice_id = _next_invoice_id(sheet)
    created_at = datetime.datetime.now().isoformat()
    due_date = (datetime.date.today() + datetime.timedelta(days=payment_terms_days)).isoformat()
    status = "ISSUED"

    sheet.append([
        invoice_id,
        sales_order_id,
        delivery_id,
        customer_id,
        customer_name,
        quantity_billed,
        unit_price,
        invoice_amount,
        payment_terms_days,
        due_date,
        status,
        notes,
        created_at,
    ])
    wb.save(MOCK_ERP_PATH)

    result = {
        "created": True,
        "invoice_id": invoice_id,
        "invoice_amount": invoice_amount,
        "due_date": due_date,
        "status": status,
    }
    record_idempotent_result(idempotency_key, "create_invoice", result)
    return result


# Newest first, same convention as every other channel's list_*() —
# the "Cash" half of Order-to-Cash finally gets a real, browsable view
# instead of only ever flashing inside one Agent Console chat session.
def list_invoices() -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "Invoices" not in wb.sheetnames:
        return []
    sheet = wb["Invoices"]
    rows = [dict(zip(INVOICE_HEADERS, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    rows.reverse()
    return rows


def find_invoice(invoice_id: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "Invoices" not in wb.sheetnames:
        return None
    sheet = wb["Invoices"]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == invoice_id:
            return dict(zip(INVOICE_HEADERS, row))
    return None


def find_invoice_by_sales_order(sales_order_id: str) -> dict | None:
    """Used by the Track Order panel — a sales order only ever gets one
    invoice in this system today (no partial/multi-invoice billing), so
    the most recent match is the real, only match."""
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "Invoices" not in wb.sheetnames:
        return None
    sheet = wb["Invoices"]
    for row in reversed(list(sheet.iter_rows(min_row=2, values_only=True))):
        if row[1] == sales_order_id:
            return dict(zip(INVOICE_HEADERS, row))
    return None


# The first real payment-tracking capability in this system — every
# credit-exposure/billing comment up to now has explicitly flagged "no
# payment tracking yet" as a known gap. Added narrowly here because
# Export's own pipeline explicitly ends in "Payment Realisation -> Closed"
# and needs a real signal to close against, not because every invoice
# flow elsewhere is being retrofitted with payment reconciliation now.
def mark_invoice_paid(invoice_id: str) -> dict:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "Invoices" not in wb.sheetnames:
        return {"updated": False, "reason": f"No invoice found with id {invoice_id}"}
    sheet = wb["Invoices"]
    status_col = INVOICE_HEADERS.index("status") + 1

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value != invoice_id:
            continue
        current_status = row[status_col - 1].value
        if current_status != "ISSUED":
            return {
                "updated": False,
                "reason": f"Cannot mark {invoice_id} paid from status {current_status} (expected ISSUED).",
            }
        sheet.cell(row=row_idx, column=status_col, value="PAID")
        wb.save(MOCK_ERP_PATH)
        return {"updated": True, "invoice_id": invoice_id, "status": "PAID"}

    return {"updated": False, "reason": f"No invoice found with id {invoice_id}"}
