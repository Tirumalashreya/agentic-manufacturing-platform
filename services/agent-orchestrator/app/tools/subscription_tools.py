# app/tools/subscription_tools.py
import datetime
import openpyxl

MOCK_ERP_PATH = "mock_erp.xlsx"

SUBSCRIPTION_HEADERS = [
    "subscription_id",
    "customer_id",
    "customer_name",
    "plan_id",
    "plan_name",
    "price",
    "billing_frequency",
    "action",
    "start_date",
    "renewal_date",
    "status",
    "created_at",
]


# Deterministic plan matching — case-insensitive substring against the
# real seeded catalog, no second LLM round-trip needed for this. Returns
# None for no match rather than guessing, same as get_product()'s
# exact-match-first philosophy.
def find_subscription_plan(requested_plan_name: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "SubscriptionPlans" not in wb.sheetnames:
        return None
    sheet = wb["SubscriptionPlans"]
    needle = requested_plan_name.strip().lower()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        plan_id, plan_name, price, billing_frequency = row
        if plan_name.strip().lower() == needle:
            return {
                "plan_id": plan_id,
                "plan_name": plan_name,
                "price": price,
                "billing_frequency": billing_frequency,
            }

    # Fall back to substring match either direction — "annual software
    # subscription" should still find "Annual Software Subscription".
    for row in sheet.iter_rows(min_row=2, values_only=True):
        plan_id, plan_name, price, billing_frequency = row
        haystack = plan_name.strip().lower()
        if needle in haystack or haystack in needle:
            return {
                "plan_id": plan_id,
                "plan_name": plan_name,
                "price": price,
                "billing_frequency": billing_frequency,
            }

    return None


def _subscriptions_sheet(wb):
    if "Subscriptions" not in wb.sheetnames:
        sheet = wb.create_sheet("Subscriptions")
        sheet.append(SUBSCRIPTION_HEADERS)
        return sheet
    return wb["Subscriptions"]


def _next_subscription_id(sheet) -> str:
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith("SUB"):
            try:
                max_num = max(max_num, int(str(value)[3:]))
            except ValueError:
                pass
    return f"SUB{max_num + 1:04d}"


def _renewal_days(billing_frequency: str) -> int:
    return 365 if billing_frequency == "ANNUAL" else 30


# Only ever appends a NEW row — RENEWAL/CANCEL act on an existing
# subscription instead (renew_subscription / cancel_subscription), they
# never call this.
def create_subscription(
    customer_id: str | None,
    customer_name: str | None,
    plan_id: str,
    plan_name: str,
    price: float,
    billing_frequency: str,
    action: str = "NEW",
) -> dict:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _subscriptions_sheet(wb)

    subscription_id = _next_subscription_id(sheet)
    created_at = datetime.datetime.now().isoformat()
    start_date = datetime.date.today().isoformat()
    renewal_date = (datetime.date.today() + datetime.timedelta(days=_renewal_days(billing_frequency))).isoformat()
    status = "ACTIVE"

    sheet.append([
        subscription_id,
        customer_id,
        customer_name,
        plan_id,
        plan_name,
        price,
        billing_frequency,
        action,
        start_date,
        renewal_date,
        status,
        created_at,
    ])
    wb.save(MOCK_ERP_PATH)

    return {
        "created": True,
        "subscription_id": subscription_id,
        "renewal_date": renewal_date,
        "status": status,
    }


# Most recent (by created_at) ACTIVE subscription for this customer+plan —
# the thing a RENEWAL/CANCEL request actually acts on. None if there's
# nothing real to renew/cancel; callers must not fabricate one.
def find_active_subscription(customer_id: str | None, plan_id: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "Subscriptions" not in wb.sheetnames:
        return None
    sheet = wb["Subscriptions"]
    best = None
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(SUBSCRIPTION_HEADERS, row))
        if record["customer_id"] != customer_id or record["plan_id"] != plan_id:
            continue
        if record["status"] != "ACTIVE":
            continue
        if best is None or record["created_at"] > best["created_at"]:
            best = record
    return best


# Extends the existing subscription's renewal_date from its current
# renewal_date (not from today) — renewing early doesn't lose the
# remaining paid-for period. Does not touch price/plan.
def renew_subscription(subscription_id: str) -> dict:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "Subscriptions" not in wb.sheetnames:
        return {"renewed": False, "reason": f"No subscription found with id {subscription_id}"}
    sheet = wb["Subscriptions"]

    status_col = SUBSCRIPTION_HEADERS.index("status") + 1
    renewal_date_col = SUBSCRIPTION_HEADERS.index("renewal_date") + 1
    billing_frequency_col = SUBSCRIPTION_HEADERS.index("billing_frequency") + 1

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value != subscription_id:
            continue
        current_status = row[status_col - 1].value
        if current_status != "ACTIVE":
            return {"renewed": False, "reason": f"{subscription_id} is not ACTIVE (status: {current_status})."}

        current_renewal_date = datetime.date.fromisoformat(row[renewal_date_col - 1].value)
        billing_frequency = row[billing_frequency_col - 1].value
        new_renewal_date = (current_renewal_date + datetime.timedelta(days=_renewal_days(billing_frequency))).isoformat()

        sheet.cell(row=row_idx, column=renewal_date_col, value=new_renewal_date)
        wb.save(MOCK_ERP_PATH)
        return {"renewed": True, "subscription_id": subscription_id, "renewal_date": new_renewal_date, "status": "ACTIVE"}

    return {"renewed": False, "reason": f"No subscription found with id {subscription_id}"}


def cancel_subscription(subscription_id: str) -> dict:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "Subscriptions" not in wb.sheetnames:
        return {"cancelled": False, "reason": f"No subscription found with id {subscription_id}"}
    sheet = wb["Subscriptions"]

    status_col = SUBSCRIPTION_HEADERS.index("status") + 1

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value != subscription_id:
            continue
        current_status = row[status_col - 1].value
        if current_status != "ACTIVE":
            return {"cancelled": False, "reason": f"{subscription_id} is not ACTIVE (status: {current_status})."}

        sheet.cell(row=row_idx, column=status_col, value="CANCELLED")
        wb.save(MOCK_ERP_PATH)
        return {"cancelled": True, "subscription_id": subscription_id, "status": "CANCELLED"}

    return {"cancelled": False, "reason": f"No subscription found with id {subscription_id}"}
