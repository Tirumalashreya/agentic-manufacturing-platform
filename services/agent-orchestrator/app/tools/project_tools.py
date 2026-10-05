# app/tools/project_tools.py
import datetime
import openpyxl

from app.erp.mock_erp_adapter import default_erp_adapter as erp
from app.tools.audit_tools import record_audit_event

MOCK_ERP_PATH = "mock_erp.xlsx"

# Contract/milestone/sign-off/retention are all 1:1 with a Project or a
# ProjectMilestone — folded directly into their rows rather than kept as
# separate near-empty sheets, same pattern already used for
# PostSalesActions (inspection fields in-row), ServiceOrders (completion
# fields in-row) and TradingOrders (every lifecycle field in-row).
PROJECT_HEADERS = [
    "project_id",
    "customer_id",
    "customer_name",
    "product_id",
    "product_name",
    "quantity",
    "project_type",
    "scope_description",
    "contract_value",
    "production_order_id",
    "retention_percent",
    "retention_held",
    "retention_released",
    "status",
    "created_at",
    "closed_at",
]

MILESTONE_HEADERS = [
    "milestone_id",
    "project_id",
    "milestone_name",
    "sequence",
    "percent_of_contract",
    "amount",
    "status",
    "signed_off_by",
    "signed_off_at",
    "invoice_id",
    "billed_amount",
]

# A fixed default split, same policy-constant pattern as
# RETURN_WINDOW_DAYS (post_sales_tools) and SERVICE_PAYMENT_TERMS_DAYS
# (service_tools) — one flat rule until there's a real need for it to
# vary per project/customer. Percentages must sum to 1.0.
MILESTONE_SPLIT = [
    ("ADVANCE", 0.30),
    ("DELIVERY", 0.40),
    ("FINAL", 0.30),
]

# Held back from every milestone billing, released as one settlement once
# the project closes — never invented per-project, same flat-until-proven
# pattern as the split above.
RETENTION_PERCENT = 0.10

_MILESTONE_TRANSITIONS = {
    "PENDING": {"IN_PROGRESS"},
    "IN_PROGRESS": {"DELIVERED"},
    "DELIVERED": {"SIGNED_OFF"},
}


def _projects_sheet(wb):
    if "Projects" not in wb.sheetnames:
        sheet = wb.create_sheet("Projects")
        sheet.append(PROJECT_HEADERS)
        return sheet
    return wb["Projects"]


def _milestones_sheet(wb):
    if "ProjectMilestones" not in wb.sheetnames:
        sheet = wb.create_sheet("ProjectMilestones")
        sheet.append(MILESTONE_HEADERS)
        return sheet
    return wb["ProjectMilestones"]


def _next_id(sheet, prefix: str) -> str:
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith(prefix):
            try:
                max_num = max(max_num, int(str(value)[len(prefix):]))
            except ValueError:
                pass
    return f"{prefix}{max_num + 1:04d}"


# Creates the Project and its milestones in one shot — a Project without
# milestones isn't a meaningful state to leave behind. contract_value may
# be None (the agent never invents one); milestone amounts stay None in
# that case, same honest "no rate on file yet" fallback create_service_order
# already uses, until a real value is available to bill against.
def create_project(
    customer_id: str | None,
    customer_name: str | None,
    product_id: str | None,
    product_name: str | None,
    quantity: int | None,
    project_type: str,
    scope_description: str,
    contract_value: float | None,
) -> dict:
    if not scope_description:
        return {
            "created": False,
            "reason": "Cannot create a project without a scope description.",
        }
    if quantity is not None and quantity <= 0:
        return {
            "created": False,
            "reason": "Cannot create a project for zero or negative quantity.",
        }

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    projects_sheet = _projects_sheet(wb)
    milestones_sheet = _milestones_sheet(wb)

    project_id = _next_id(projects_sheet, "PRJ")
    created_at = datetime.datetime.now().isoformat()

    projects_sheet.append([
        project_id,
        customer_id,
        customer_name,
        product_id,
        product_name,
        quantity,
        project_type,
        scope_description,
        contract_value,
        None,  # production_order_id — filled in once the workflow creates one
        RETENTION_PERCENT,
        0.0,
        False,
        "AWARDED",
        created_at,
        None,
    ])

    milestones = []
    for sequence, (name, percent) in enumerate(MILESTONE_SPLIT, start=1):
        milestone_id = _next_id(milestones_sheet, "MS")
        amount = round(contract_value * percent, 2) if contract_value is not None else None
        milestones_sheet.append([
            milestone_id,
            project_id,
            name,
            sequence,
            percent,
            amount,
            "PENDING",
            None,
            None,
            None,
            None,
        ])
        milestones.append({"milestone_id": milestone_id, "milestone_name": name, "amount": amount})

    wb.save(MOCK_ERP_PATH)
    record_audit_event("Project", project_id, "PROJECT_CREATED", {
        "project_type": project_type, "contract_value": contract_value,
    })

    return {
        "created": True,
        "project_id": project_id,
        "status": "AWARDED",
        "milestones": milestones,
    }


# Full history for the Project Sales channel view — every project
# regardless of status, newest first.
def list_projects() -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "Projects" not in wb.sheetnames:
        return []
    sheet = wb["Projects"]
    rows = [dict(zip(PROJECT_HEADERS, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    rows.reverse()
    return rows


def find_project(project_id: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "Projects" not in wb.sheetnames:
        return None
    sheet = wb["Projects"]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == project_id:
            return dict(zip(PROJECT_HEADERS, row))
    return None


def find_project_milestones(project_id: str) -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ProjectMilestones" not in wb.sheetnames:
        return []
    sheet = wb["ProjectMilestones"]
    return [
        dict(zip(MILESTONE_HEADERS, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[1] == project_id
    ]


# The Active Workflows queue's source for "which milestones need a human
# to move them forward" — anything short of BILLED. A project's own
# status only ever reaches AWARDED or CLOSED (see create_project/
# close_project); the milestone is what actually carries the in-progress
# state a queue needs to surface.
def list_open_milestones() -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ProjectMilestones" not in wb.sheetnames:
        return []
    sheet = wb["ProjectMilestones"]
    return [
        dict(zip(MILESTONE_HEADERS, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[6] != "BILLED"
    ]


# Records a real production order once the workflow has created one — kept
# separate from create_project() because material assessment/production
# creation happens in graph.py (reusing the exact same helpers MTO/CTO
# already use), not duplicated here.
def attach_production_order(project_id: str, production_order_id: str) -> None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = wb["Projects"]
    col = PROJECT_HEADERS.index("production_order_id") + 1
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value == project_id:
            sheet.cell(row=row_idx, column=col, value=production_order_id)
            wb.save(MOCK_ERP_PATH)
            return


# PENDING -> IN_PROGRESS -> DELIVERED -> SIGNED_OFF. BILLED is deliberately
# not reachable here — see bill_milestone(), which also moves money, same
# split as update_service_order_status()/complete_service_order() and
# update_trading_order_status()/dispatch_trading_order().
def advance_milestone_status(milestone_id: str, new_status: str, signed_off_by: str | None = None) -> dict:
    if new_status not in {"IN_PROGRESS", "DELIVERED", "SIGNED_OFF"}:
        return {"updated": False, "reason": f"Invalid status: {new_status}"}

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ProjectMilestones" not in wb.sheetnames:
        return {"updated": False, "reason": f"No milestone found with id {milestone_id}"}
    sheet = wb["ProjectMilestones"]

    status_col = MILESTONE_HEADERS.index("status") + 1
    signed_off_by_col = MILESTONE_HEADERS.index("signed_off_by") + 1
    signed_off_at_col = MILESTONE_HEADERS.index("signed_off_at") + 1

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value != milestone_id:
            continue
        current_status = row[status_col - 1].value
        allowed = _MILESTONE_TRANSITIONS.get(current_status, set())
        if new_status not in allowed:
            return {
                "updated": False,
                "reason": (
                    f"Cannot move {milestone_id} from {current_status} to {new_status}. "
                    f"Allowed next status(es): {sorted(allowed) or 'none'}."
                ),
            }

        sheet.cell(row=row_idx, column=status_col, value=new_status)
        if new_status == "SIGNED_OFF":
            sheet.cell(row=row_idx, column=signed_off_by_col, value=signed_off_by or "SYSTEM")
            sheet.cell(row=row_idx, column=signed_off_at_col, value=datetime.datetime.now().isoformat())
        wb.save(MOCK_ERP_PATH)
        record_audit_event("ProjectMilestone", milestone_id, f"MILESTONE_{new_status}", {"signed_off_by": signed_off_by})
        return {"updated": True, "milestone_id": milestone_id, "status": new_status}

    return {"updated": False, "reason": f"No milestone found with id {milestone_id}"}


# The billing step. Only a SIGNED_OFF milestone can be billed. Bills
# amount net of retention — the held-back portion accumulates on the
# Project row and is only released as one settlement at close_project(),
# never per-milestone. A milestone with no amount (contract_value was
# never stated) still moves to BILLED — the work happened — but produces
# no invoice, same honest "no rate on file" fallback as
# complete_service_order().
def bill_milestone(milestone_id: str) -> dict:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ProjectMilestones" not in wb.sheetnames:
        return {"billed": False, "reason": f"No milestone found with id {milestone_id}"}
    milestones_sheet = wb["ProjectMilestones"]

    status_col = MILESTONE_HEADERS.index("status") + 1
    project_id_col = MILESTONE_HEADERS.index("project_id") + 1
    name_col = MILESTONE_HEADERS.index("milestone_name") + 1
    amount_col = MILESTONE_HEADERS.index("amount") + 1
    invoice_id_col = MILESTONE_HEADERS.index("invoice_id") + 1
    billed_amount_col = MILESTONE_HEADERS.index("billed_amount") + 1

    for row_idx, row in enumerate(milestones_sheet.iter_rows(min_row=2), start=2):
        if row[0].value != milestone_id:
            continue
        current_status = row[status_col - 1].value
        if current_status != "SIGNED_OFF":
            return {
                "billed": False,
                "reason": f"{milestone_id} must be SIGNED_OFF before it can be billed (status: {current_status}).",
            }

        project_id = row[project_id_col - 1].value
        milestone_name = row[name_col - 1].value
        amount = row[amount_col - 1].value

        project = find_project(project_id)
        invoice_id = None
        billed_amount = None
        retained_amount = 0.0

        if amount is not None:
            retention_percent = project["retention_percent"] or 0.0
            retained_amount = round(amount * retention_percent, 2)
            billed_amount = round(amount - retained_amount, 2)
            invoice = erp.create_invoice(
                sales_order_id=None,
                delivery_id=None,
                customer_id=project["customer_id"],
                customer_name=project["customer_name"],
                quantity_billed=1,
                unit_price=billed_amount,
                payment_terms_days=30,
                notes=f"Project {project_id} milestone {milestone_name} ({milestone_id}) signed off and billed, net of {retention_percent:.0%} retention.",
            )
            if invoice["created"]:
                invoice_id = invoice["invoice_id"]

        milestones_sheet.cell(row=row_idx, column=status_col, value="BILLED")
        milestones_sheet.cell(row=row_idx, column=invoice_id_col, value=invoice_id)
        milestones_sheet.cell(row=row_idx, column=billed_amount_col, value=billed_amount)

        if retained_amount:
            projects_sheet = wb["Projects"]
            held_col = PROJECT_HEADERS.index("retention_held") + 1
            for p_row_idx, p_row in enumerate(projects_sheet.iter_rows(min_row=2), start=2):
                if p_row[0].value == project_id:
                    current_held = p_row[held_col - 1].value or 0.0
                    projects_sheet.cell(row=p_row_idx, column=held_col, value=round(current_held + retained_amount, 2))
                    break

        wb.save(MOCK_ERP_PATH)
        record_audit_event("ProjectMilestone", milestone_id, "MILESTONE_BILLED", {
            "invoice_id": invoice_id, "billed_amount": billed_amount, "retained_amount": retained_amount,
        })
        return {
            "billed": True,
            "milestone_id": milestone_id,
            "invoice_id": invoice_id,
            "billed_amount": billed_amount,
            "retained_amount": retained_amount,
        }

    return {"billed": False, "reason": f"No milestone found with id {milestone_id}"}


# Only closable once every milestone is BILLED — a project can't close
# with unresolved work. Retention is released as one final settlement
# invoice, matching the "held back until the very end" semantics
# real construction/EPC contracts use, rather than trickling it out
# per-milestone.
def close_project(project_id: str) -> dict:
    project = find_project(project_id)
    if project is None:
        return {"closed": False, "reason": f"No project found with id {project_id}"}
    if project["status"] == "CLOSED":
        return {"closed": False, "reason": f"{project_id} is already CLOSED."}

    milestones = find_project_milestones(project_id)
    unbilled = [m for m in milestones if m["status"] != "BILLED"]
    if unbilled:
        return {
            "closed": False,
            "reason": (
                f"Cannot close {project_id} — {len(unbilled)} milestone(s) not yet BILLED: "
                f"{', '.join(m['milestone_id'] for m in unbilled)}."
            ),
        }

    retention_held = project["retention_held"] or 0.0
    invoice_id = None
    if retention_held > 0:
        invoice = erp.create_invoice(
            sales_order_id=None,
            delivery_id=None,
            customer_id=project["customer_id"],
            customer_name=project["customer_name"],
            quantity_billed=1,
            unit_price=retention_held,
            payment_terms_days=30,
            notes=f"Project {project_id} retention release on close.",
        )
        if invoice["created"]:
            invoice_id = invoice["invoice_id"]

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = wb["Projects"]
    status_col = PROJECT_HEADERS.index("status") + 1
    closed_at_col = PROJECT_HEADERS.index("closed_at") + 1
    released_col = PROJECT_HEADERS.index("retention_released") + 1

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value != project_id:
            continue
        sheet.cell(row=row_idx, column=status_col, value="CLOSED")
        sheet.cell(row=row_idx, column=closed_at_col, value=datetime.datetime.now().isoformat())
        sheet.cell(row=row_idx, column=released_col, value=True)
        wb.save(MOCK_ERP_PATH)
        break

    record_audit_event("Project", project_id, "PROJECT_CLOSED", {
        "retention_released": retention_held, "invoice_id": invoice_id,
    })
    return {
        "closed": True,
        "project_id": project_id,
        "retention_released": retention_held,
        "invoice_id": invoice_id,
    }
