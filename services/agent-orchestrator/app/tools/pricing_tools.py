# app/tools/pricing_tools.py
import openpyxl

MOCK_ERP_PATH = "mock_erp.xlsx"

# Policy DATA lives in the DiscountPolicy sheet, not here — this is only
# the seed used the first time that sheet doesn't exist yet. Changing the
# discount rules means editing sheet rows, not this file or the
# negotiation agent.
DISCOUNT_POLICY_HEADERS = [
    "rule_type",
    "min_threshold",
    "max_threshold",
    "discount_percent",
    "label",
]

_DEFAULT_DISCOUNT_POLICY_ROWS = [
    # rule_type              min_threshold  max_threshold  discount_percent  label
    ["PRODUCT_BASE",         0,             5000,          3,   "Tier A - Accessories"],
    ["PRODUCT_BASE",         5000,          50000,         5,   "Tier B - Standard"],
    ["PRODUCT_BASE",         50000,         None,          8,   "Tier C - High-value / Enterprise"],
    ["QUANTITY_BONUS",       0,             20,            0,   "Small order"],
    ["QUANTITY_BONUS",       20,            100,           2,   "Medium order"],
    ["QUANTITY_BONUS",       100,           None,          4,   "Bulk order"],
    ["CUSTOMER_TIER_BONUS",  0,             200000,        0,   "Standard tier"],
    ["CUSTOMER_TIER_BONUS",  200000,        500000,        2,   "Preferred tier"],
    ["CUSTOMER_TIER_BONUS",  500000,        None,          4,   "Premium tier"],
    ["OVERALL_CAP",          None,          None,          15,  "Overall ceiling"],
]


def _discount_policy_sheet(wb):
    if "DiscountPolicy" not in wb.sheetnames:
        sheet = wb.create_sheet("DiscountPolicy")
        sheet.append(DISCOUNT_POLICY_HEADERS)
        for row in _DEFAULT_DISCOUNT_POLICY_ROWS:
            sheet.append(row)
        wb.save(MOCK_ERP_PATH)
        return sheet
    return wb["DiscountPolicy"]


def _load_discount_policy() -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _discount_policy_sheet(wb)
    return [
        dict(zip(DISCOUNT_POLICY_HEADERS, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
    ]


def _bracket_percent(rules: list[dict], rule_type: str, value: float) -> float:
    for rule in rules:
        if rule["rule_type"] != rule_type:
            continue
        lower = rule["min_threshold"]
        upper = rule["max_threshold"]
        if lower is not None and value < lower:
            continue
        if upper is not None and value >= upper:
            continue
        return rule["discount_percent"]
    return 0.0


# Product/quantity/customer-tier -> the maximum discount this specific
# negotiation is allowed to auto-approve. Pure lookup + arithmetic over
# sheet data — no hardcoded numbers here, no LLM involvement.
def compute_max_discount(
    product_price: float,
    quantity: int,
    customer_credit_limit: float,
) -> dict:
    rules = _load_discount_policy()

    product_base = _bracket_percent(rules, "PRODUCT_BASE", product_price)
    quantity_bonus = _bracket_percent(rules, "QUANTITY_BONUS", quantity)
    tier_bonus = _bracket_percent(rules, "CUSTOMER_TIER_BONUS", customer_credit_limit)
    cap = next(r["discount_percent"] for r in rules if r["rule_type"] == "OVERALL_CAP")

    raw_total = product_base + quantity_bonus + tier_bonus
    max_discount_percent = min(raw_total, cap)

    return {
        "product_base_percent": product_base,
        "quantity_bonus_percent": quantity_bonus,
        "customer_tier_bonus_percent": tier_bonus,
        "raw_total_percent": raw_total,
        "cap_percent": cap,
        "max_discount_percent": max_discount_percent,
    }


def evaluate_negotiation(
    original_price: float,
    quantity: int,
    customer_credit_limit: float,
    requested_discount_percent: float | None,
    requested_price: float | None,
) -> dict:
    if original_price <= 0:
        return {
            "evaluated": False,
            "reason": "Original price must be greater than zero.",
        }

    if requested_discount_percent is None and requested_price is None:
        return {
            "evaluated": False,
            "reason": "No discount percentage or target price was specified.",
        }

    # Deterministic conversion between the two forms — never the model's job.
    if requested_discount_percent is None:
        requested_discount_percent = (original_price - requested_price) / original_price * 100
    if requested_price is None:
        requested_price = original_price * (1 - requested_discount_percent / 100)

    policy = compute_max_discount(original_price, quantity, customer_credit_limit)
    max_auto_approved_percent = policy["max_discount_percent"]
    within_authority = requested_discount_percent <= max_auto_approved_percent
    status = "REVISED" if within_authority else "REQUIRES_APPROVAL"

    return {
        "evaluated": True,
        "original_price": original_price,
        "requested_price": round(requested_price, 2),
        "requested_discount_percent": round(requested_discount_percent, 2),
        "max_auto_approved_percent": max_auto_approved_percent,
        "policy_breakdown": policy,
        "within_authority": within_authority,
        "revised_price": round(requested_price, 2) if within_authority else None,
        "status": status,
    }
