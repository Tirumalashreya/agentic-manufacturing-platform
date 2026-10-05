# app/tools/audit_tools.py
import datetime
import json
import openpyxl

MOCK_ERP_PATH = "mock_erp.xlsx"

AUDIT_EVENT_HEADERS = ["event_id", "entity_type", "entity_id", "event_type", "details", "created_at"]

# One shared, append-only event log, not a bespoke history table per
# entity. Wired into the multi-stage workflows this session built where a
# real timeline is actually valuable (Project milestones, Export,
# Trading, Service, Consignment, Financial, Post-Sales) — not every
# status flip everywhere, per the "add reliability where it matters, not
# because it's on a checklist" rule.


def _sheet(wb):
    if "AuditEvents" not in wb.sheetnames:
        sheet = wb.create_sheet("AuditEvents")
        sheet.append(AUDIT_EVENT_HEADERS)
        return sheet
    return wb["AuditEvents"]


def _next_event_id(sheet) -> str:
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith("EVT"):
            try:
                max_num = max(max_num, int(str(value)[3:]))
            except ValueError:
                pass
    return f"EVT{max_num + 1:06d}"


def record_audit_event(
    entity_type: str,
    entity_id: str | None,
    event_type: str,
    details: dict | None = None,
) -> None:
    if not entity_id:
        return
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _sheet(wb)
    event_id = _next_event_id(sheet)
    sheet.append([
        event_id,
        entity_type,
        entity_id,
        event_type,
        json.dumps(details or {}),
        datetime.datetime.now().isoformat(),
    ])
    wb.save(MOCK_ERP_PATH)


def get_audit_trail(entity_type: str, entity_id: str) -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "AuditEvents" not in wb.sheetnames:
        return []
    sheet = wb["AuditEvents"]
    trail = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        event_id, row_entity_type, row_entity_id, event_type, details_json, created_at = row
        if row_entity_type == entity_type and row_entity_id == entity_id:
            trail.append({
                "event_id": event_id,
                "entity_type": row_entity_type,
                "entity_id": row_entity_id,
                "event_type": event_type,
                "details": json.loads(details_json) if details_json else {},
                "created_at": created_at,
            })
    return trail


# The cross-entity feed a real Audit Trail screen needs — get_audit_trail()
# is scoped to one already-known entity, which a global feed doesn't have.
# Newest first, based on insertion order (this sheet is only ever
# appended to), same "most recent activity" framing as any real audit log.
def list_recent_audit_events(limit: int = 100) -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "AuditEvents" not in wb.sheetnames:
        return []
    sheet = wb["AuditEvents"]
    events = [
        {
            "event_id": event_id,
            "entity_type": entity_type,
            "entity_id": entity_id,
            "event_type": event_type,
            "details": json.loads(details_json) if details_json else {},
            "created_at": created_at,
        }
        for event_id, entity_type, entity_id, event_type, details_json, created_at
        in sheet.iter_rows(min_row=2, values_only=True)
    ]
    events.reverse()
    return events[:limit]
