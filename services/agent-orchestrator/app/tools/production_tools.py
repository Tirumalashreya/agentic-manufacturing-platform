# app/tools/production_tools.py
import datetime
import openpyxl

MOCK_ERP_PATH = "mock_erp.xlsx"

PRODUCTION_ORDER_HEADERS = [
    "production_order_id",
    "sales_order_id",
    "product_id",
    "product_name",
    "quantity",
    "requires_configuration",
    "configuration_notes",
    "target_completion_date",
    "status",
    "created_at",
]

RESERVATION_HEADERS = [
    "reservation_id",
    "sales_order_id",
    "product_id",
    "product_name",
    "plant",
    "quantity",
    "created_at",
]


def _production_orders_sheet(wb):
    if "ProductionOrders" not in wb.sheetnames:
        sheet = wb.create_sheet("ProductionOrders")
        sheet.append(PRODUCTION_ORDER_HEADERS)
        return sheet
    return wb["ProductionOrders"]


def compute_target_completion_date(lead_time_days: int | None) -> str:
    days = lead_time_days if lead_time_days is not None else 30
    return (datetime.date.today() + datetime.timedelta(days=days)).isoformat()


def _next_production_order_id(sheet) -> str:
    # Based on the max existing numeric suffix, not row count — robust to
    # any row ever being deleted.
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith("PO"):
            try:
                max_num = max(max_num, int(str(value)[2:]))
            except ValueError:
                pass
    return f"PO{max_num + 1:04d}"


# MTO/CTO/HYBRID(production leg) all call this. The Configuration Agent
# doesn't exist yet, so a CTO order is recorded honestly as "requires
# configuration" rather than pretending the request was captured.
def create_production_order(
    sales_order_id: str,
    product_id: str,
    product_name: str,
    quantity: int,
    lead_time_days: int | None,
    requires_configuration: bool = False,
    configuration_notes: str | None = None,
    initial_status: str = "PLANNED",
) -> dict:
    if quantity is None or quantity <= 0:
        return {
            "created": False,
            "reason": "Cannot create a production order for zero or negative quantity.",
        }

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _production_orders_sheet(wb)

    production_order_id = _next_production_order_id(sheet)
    created_at = datetime.datetime.now().isoformat()
    target_completion_date = compute_target_completion_date(lead_time_days)
    status = initial_status

    sheet.append([
        production_order_id,
        sales_order_id,
        product_id,
        product_name,
        quantity,
        requires_configuration,
        configuration_notes,
        target_completion_date,
        status,
        created_at,
    ])
    wb.save(MOCK_ERP_PATH)

    return {
        "created": True,
        "production_order_id": production_order_id,
        "target_completion_date": target_completion_date,
        "status": status,
    }


def _reservations_sheet(wb):
    if "StockReservations" not in wb.sheetnames:
        sheet = wb.create_sheet("StockReservations")
        sheet.append(RESERVATION_HEADERS)
        return sheet
    return wb["StockReservations"]


def _next_reservation_id(sheet) -> str:
    # Based on the max existing numeric suffix, not row count — robust to
    # any row ever being deleted.
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith("RSV"):
            try:
                max_num = max(max_num, int(str(value)[3:]))
            except ValueError:
                pass
    return f"RSV{max_num + 1:04d}"


# MTS/HYBRID(stock leg) call this. Actually decrements Inventory — this is
# the one tool in this file that mutates existing data rather than just
# appending a new record, so it guards against a torn state: if the real,
# current Inventory can't cover the request, nothing is written at all.
def reserve_stock(
    sales_order_id: str,
    product_id: str,
    product_name: str,
    quantity: int,
) -> dict:
    if quantity is None or quantity <= 0:
        return {
            "reserved": False,
            "reason": "Cannot reserve zero or negative quantity.",
        }

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    inventory = wb["Inventory"]

    remaining = quantity
    allocations = []
    touched_cells = []
    for row in inventory.iter_rows(min_row=2):
        if remaining <= 0:
            break
        pid_cell, plant_cell, qty_cell = row
        if pid_cell.value != product_id:
            continue
        available = qty_cell.value or 0
        if available <= 0:
            continue
        take = min(available, remaining)
        touched_cells.append((qty_cell, available - take))
        allocations.append({"plant": plant_cell.value, "quantity": take})
        remaining -= take

    if remaining > 0:
        # Nothing has been saved yet — the loaded workbook is discarded
        # unsaved, so no partial mutation ever reaches the file.
        return {
            "reserved": False,
            "reason": f"Insufficient stock: short by {remaining} units.",
        }

    for qty_cell, new_value in touched_cells:
        qty_cell.value = new_value

    reservations_sheet = _reservations_sheet(wb)
    reservation_id = _next_reservation_id(reservations_sheet)
    created_at = datetime.datetime.now().isoformat()
    for alloc in allocations:
        reservations_sheet.append([
            reservation_id,
            sales_order_id,
            product_id,
            product_name,
            alloc["plant"],
            alloc["quantity"],
            created_at,
        ])
    wb.save(MOCK_ERP_PATH)

    return {
        "reserved": True,
        "reservation_id": reservation_id,
        "allocations": allocations,
    }
