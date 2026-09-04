# app/tools/product_tools.py
import openpyxl

MOCK_ERP_PATH = "mock_erp.xlsx"


def get_product(product_name: str | None) -> dict | None:
    if product_name is None:
        # A genuine "no product at all" request — a standalone financial
        # adjustment, for example. Not every request names a product.
        return None

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = wb["Product Master"]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        product_id, name, price, configurable, can_produce, lead_time_days = row
        if name.lower() == product_name.lower():
            return {
                "product_id": product_id,
                "product_name": name,
                "price": price,
                "configurable": bool(configurable),
                "can_produce": bool(can_produce),
                "lead_time_days": lead_time_days,
            }
    return None


def get_all_product_names() -> list[str]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = wb["Product Master"]
    return [row[0] for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True)]
