# app/tools/sales_tools.py
import datetime
import openpyxl

from app.tools.idempotency_tools import get_idempotent_result, record_idempotent_result

MOCK_ERP_PATH = "mock_erp.xlsx"

QUOTATION_HEADERS = [
    "quotation_id",
    "customer_id",
    "customer_name",
    "product_id",
    "product_name",
    "quantity",
    "quoted_price",
    "total_amount",
    "valid_for_days",
    "valid_until",
    "status",
    "notes",
    "created_at",
]

SALES_ORDER_HEADERS = [
    "sales_order_id",
    "customer_id",
    "customer_name",
    "product_id",
    "product_name",
    "quantity",
    "unit_price",
    "total_amount",
    "requested_deadline",
    "status",
    "created_at",
]


def _quotations_sheet(wb):
    if "Quotations" not in wb.sheetnames:
        sheet = wb.create_sheet("Quotations")
        sheet.append(QUOTATION_HEADERS)
        return sheet
    return wb["Quotations"]


def _next_id(sheet, prefix: str) -> str:
    # Based on the max existing numeric suffix, not row count — robust to
    # any row ever being deleted (row-count-based IDs can collide with a
    # still-existing higher-numbered row after a deletion).
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith(prefix):
            try:
                max_num = max(max_num, int(str(value)[len(prefix):]))
            except ValueError:
                pass
    return f"{prefix}{max_num + 1:04d}"


def _next_quotation_id(sheet) -> str:
    return _next_id(sheet, "Q")


def create_quotation(
    customer_id: str | None,
    customer_name: str | None,
    product_id: str | None,
    product_name: str | None,
    quantity: int | None,
    quoted_price: float | None,
    valid_for_days: int,
    notes: str,
) -> dict:
    # Deterministic guard — a quotation with no real product, no price, or
    # no sensible quantity isn't a business transaction worth recording.
    if (
        product_id is None
        or quoted_price is None
        or quantity is None
        or quantity <= 0
    ):
        return {
            "created": False,
            "reason": "Cannot create a quotation without a valid product, price, and quantity.",
        }

    # Computed here, not accepted as a caller-supplied value — a total
    # cannot silently drift from quoted_price * quantity.
    total_amount = round(quoted_price * quantity, 2)

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _quotations_sheet(wb)

    quotation_id = _next_quotation_id(sheet)
    created_at = datetime.datetime.now().isoformat()
    valid_until = (datetime.date.today() + datetime.timedelta(days=valid_for_days)).isoformat()
    status = "OPEN"

    sheet.append([
        quotation_id,
        customer_id,
        customer_name,
        product_id,
        product_name,
        quantity,
        quoted_price,
        total_amount,
        valid_for_days,
        valid_until,
        status,
        notes,
        created_at,
    ])
    wb.save(MOCK_ERP_PATH)

    return {
        "created": True,
        "quotation_id": quotation_id,
        "total_amount": total_amount,
        "valid_until": valid_until,
        "status": status,
    }


def _sales_orders_sheet(wb):
    if "SalesOrders" not in wb.sheetnames:
        sheet = wb.create_sheet("SalesOrders")
        sheet.append(SALES_ORDER_HEADERS)
        return sheet
    return wb["SalesOrders"]


def _next_sales_order_id(sheet) -> str:
    return _next_id(sheet, "SO")


def create_sales_order(
    customer_id: str | None,
    customer_name: str | None,
    customer_active: bool,
    product_id: str | None,
    product_name: str | None,
    quantity: int | None,
    unit_price: float | None,
    requested_deadline: str | None,
    idempotency_key: str | None = None,
) -> dict:
    cached = get_idempotent_result(idempotency_key, "create_sales_order")
    if cached is not None:
        return cached

    # A Sales Order is a real commitment, not a pre-sales quote to a
    # prospect — unlike create_quotation(), it requires an actual,
    # active registered customer, not just a valid product/price.
    if customer_id is None or not customer_active:
        return {
            "created": False,
            "reason": "Cannot create a sales order for an unknown or inactive customer.",
        }
    if product_id is None or unit_price is None:
        return {
            "created": False,
            "reason": "Cannot create a sales order without a valid product and price.",
        }
    if quantity is None or quantity <= 0:
        return {
            "created": False,
            "reason": "Cannot create a sales order without a valid quantity.",
        }

    total_amount = round(unit_price * quantity, 2)

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _sales_orders_sheet(wb)

    sales_order_id = _next_sales_order_id(sheet)
    created_at = datetime.datetime.now().isoformat()
    status = "CONFIRMED"

    sheet.append([
        sales_order_id,
        customer_id,
        customer_name,
        product_id,
        product_name,
        quantity,
        unit_price,
        total_amount,
        requested_deadline,
        status,
        created_at,
    ])
    wb.save(MOCK_ERP_PATH)

    result = {
        "created": True,
        "sales_order_id": sales_order_id,
        "total_amount": total_amount,
        "status": status,
    }
    record_idempotent_result(idempotency_key, "create_sales_order", result)
    return result


# The persisted SalesOrders row must reflect reality — a credit hold means
# the real record says CREDIT_BLOCKED, not just LangGraph state saying
# ON_CREDIT_HOLD while the sheet still claims CONFIRMED.
def update_sales_order_status(sales_order_id: str, new_status: str) -> dict:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _sales_orders_sheet(wb)

    status_col = SALES_ORDER_HEADERS.index("status") + 1
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value == sales_order_id:
            sheet.cell(row=row_idx, column=status_col, value=new_status)
            wb.save(MOCK_ERP_PATH)
            return {"updated": True, "sales_order_id": sales_order_id, "status": new_status}

    return {"updated": False, "reason": f"No sales order found with id {sales_order_id}"}


# The Domestic Sales channel view's full history, not just the ones stuck
# on credit hold — newest first, same "most recent activity" framing as
# list_recent_audit_events().
def list_sales_orders() -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "SalesOrders" not in wb.sheetnames:
        return []
    sheet = wb["SalesOrders"]
    rows = [dict(zip(SALES_ORDER_HEADERS, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    rows.reverse()
    return rows


# Used by the new Ledger order-detail page (Phase 1, agreed with the
# user) to resolve one real Sales Order's own fields — no new sheet, no
# new writes, just a single-row lookup on the same real data
# list_sales_orders already reads.
def find_sales_order(sales_order_id: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "SalesOrders" not in wb.sheetnames:
        return None
    sheet = wb["SalesOrders"]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == sales_order_id:
            return dict(zip(SALES_ORDER_HEADERS, row))
    return None


# The only status a Sales Order can be stuck in — CONFIRMED orders already
# ran fulfillment synchronously in the same graph turn (see graph.py's
# DIRECT_SALES path), so CREDIT_BLOCKED is the one real "needs a human"
# state this sheet's own status column ever reaches.
def list_credit_blocked_orders() -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "SalesOrders" not in wb.sheetnames:
        return []
    sheet = wb["SalesOrders"]
    return [
        dict(zip(SALES_ORDER_HEADERS, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[9] == "CREDIT_BLOCKED"
    ]
