# app/tools/customer_tools.py
import openpyxl

from app.tools.audit_tools import record_audit_event

MOCK_ERP_PATH = "mock_erp.xlsx"

# A brand-new customer has no order history, so there's nothing real to
# base a credit limit on — this is a stated business policy (every new
# account starts here, same limit for everyone, until someone reviews
# and raises it), not a guessed number. Kept as a named constant, same
# pattern as RETURN_WINDOW_DAYS/SERVICE_PAYMENT_TERMS_DAYS elsewhere.
NEW_CUSTOMER_STARTER_CREDIT_LIMIT = 50_000.0


def get_customer(customer_name: str | None) -> dict | None:
    # A genuine "no customer named at all" request (e.g. an implicit
    # reference intent_agent couldn't resolve to a name) — same honest
    # None-in-None-out contract get_product() already has, not an error.
    if customer_name is None:
        return None

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = wb["Customer Master"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))

    for customer_id, name, active, credit_limit in rows:
        if name.lower() == customer_name.lower():
            return {
                "customer_id": customer_id,
                "customer_name": name,
                "active": bool(active),
                "credit_limit": credit_limit,
            }

    # Fall back to substring match either direction — same deterministic
    # pattern as find_subscription_plan()/find_service_rate() — so a
    # shorthand like "ABC" still resolves to "ABC Corp". No LLM round-trip
    # needed; unlike product names, no real customer name here is a
    # substring of another, so this stays unambiguous.
    needle = customer_name.strip().lower()
    for customer_id, name, active, credit_limit in rows:
        haystack = name.strip().lower()
        if needle in haystack or haystack in needle:
            return {
                "customer_id": customer_id,
                "customer_name": name,
                "active": bool(active),
                "credit_limit": credit_limit,
            }

    return None


def _next_customer_id(sheet) -> str:
    # Same "max existing numeric suffix" pattern used everywhere else in
    # this codebase (sales_tools._next_id, export_tools, ...) — robust to
    # row deletion, not dependent on row count.
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith("C"):
            try:
                max_num = max(max_num, int(str(value)[1:]))
            except ValueError:
                pass
    return f"C{max_num + 1:03d}"


# Auto-provisioning a new customer (agreed with the user): every request
# that names a customer with no existing record gets one created on the
# spot — active, with the flat starter credit limit above — rather than
# dead-ending the request. This is the ONE place a Customer Master row
# gets created outside the original seed data, so it's the one place
# that decision needs explaining, not silently guessed at downstream.
def create_customer(customer_name: str) -> dict:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = wb["Customer Master"]
    customer_id = _next_customer_id(sheet)
    sheet.append([customer_id, customer_name, True, NEW_CUSTOMER_STARTER_CREDIT_LIMIT])
    wb.save(MOCK_ERP_PATH)

    record_audit_event(
        "Customer", customer_id, "CUSTOMER_AUTO_PROVISIONED",
        {"customer_name": customer_name, "starter_credit_limit": NEW_CUSTOMER_STARTER_CREDIT_LIMIT},
    )

    return {
        "customer_id": customer_id,
        "customer_name": customer_name,
        "active": True,
        "credit_limit": NEW_CUSTOMER_STARTER_CREDIT_LIMIT,
    }


def list_customers() -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = wb["Customer Master"]
    return [
        {"customer_id": customer_id, "customer_name": name, "active": bool(active), "credit_limit": credit_limit}
        for customer_id, name, active, credit_limit in sheet.iter_rows(min_row=2, values_only=True)
    ]
