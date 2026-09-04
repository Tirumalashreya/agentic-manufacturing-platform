# app/tools/procurement_tools.py
import datetime
import openpyxl

MOCK_ERP_PATH = "mock_erp.xlsx"

PURCHASE_ORDER_HEADERS = [
    "purchase_order_id",
    "production_order_id",
    "component_id",
    "component_name",
    "vendor_id",
    "vendor_name",
    "quantity",
    "unit_cost",
    "total_cost",
    "expected_arrival",
    "status",
    "created_at",
]


# Deterministic — reads the real BOM against the real component inventory.
# Returns has_bom=False for any product with no BOM row at all (most
# products), so callers can skip procurement honestly rather than
# pretending a shortage check happened when there's nothing to check.
def check_material_availability(product_id: str, quantity: int) -> dict:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "BOM" not in wb.sheetnames:
        return {"has_bom": False, "components": []}

    bom_sheet = wb["BOM"]
    requirements = []
    for row in bom_sheet.iter_rows(min_row=2, values_only=True):
        _bom_id, pid, component_id, component_name, qty_per_unit = row
        if pid == product_id:
            requirements.append({
                "component_id": component_id,
                "component_name": component_name,
                "required": qty_per_unit * quantity,
            })

    if not requirements:
        return {"has_bom": False, "components": []}

    inv_sheet = wb["ComponentInventory"]
    available_by_component: dict[str, int] = {}
    for row in inv_sheet.iter_rows(min_row=2, values_only=True):
        component_id, _name, _plant, stock_qty = row
        available_by_component[component_id] = available_by_component.get(component_id, 0) + (stock_qty or 0)

    components = []
    for req in requirements:
        available = available_by_component.get(req["component_id"], 0)
        shortage = max(0, req["required"] - available)
        components.append({**req, "available": available, "shortage": shortage})

    return {"has_bom": True, "components": components}


def get_vendors_for_component(component_id: str) -> list[dict]:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = wb["Vendors"]
    vendors = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        vendor_id, vendor_name, cid, lead_time_days, unit_cost = row
        if cid == component_id:
            vendors.append({
                "vendor_id": vendor_id,
                "vendor_name": vendor_name,
                "lead_time_days": lead_time_days,
                "unit_cost": unit_cost,
            })
    return vendors


def _purchase_orders_sheet(wb):
    if "PurchaseOrders" not in wb.sheetnames:
        sheet = wb.create_sheet("PurchaseOrders")
        sheet.append(PURCHASE_ORDER_HEADERS)
        return sheet
    return wb["PurchaseOrders"]


def _next_purchase_order_id(sheet) -> str:
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith("PUR"):
            try:
                max_num = max(max_num, int(str(value)[3:]))
            except ValueError:
                pass
    return f"PUR{max_num + 1:04d}"


def create_purchase_order(
    production_order_id: str,
    component_id: str,
    component_name: str,
    vendor_id: str,
    vendor_name: str,
    quantity: int,
    unit_cost: float,
    lead_time_days: int,
) -> dict:
    if quantity is None or quantity <= 0:
        return {
            "created": False,
            "reason": "Cannot create a purchase order for zero or negative quantity.",
        }

    total_cost = round(unit_cost * quantity, 2)

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _purchase_orders_sheet(wb)

    purchase_order_id = _next_purchase_order_id(sheet)
    created_at = datetime.datetime.now().isoformat()
    expected_arrival = (datetime.date.today() + datetime.timedelta(days=lead_time_days)).isoformat()
    status = "ORDERED"

    sheet.append([
        purchase_order_id,
        production_order_id,
        component_id,
        component_name,
        vendor_id,
        vendor_name,
        quantity,
        unit_cost,
        total_cost,
        expected_arrival,
        status,
        created_at,
    ])
    wb.save(MOCK_ERP_PATH)

    return {
        "created": True,
        "purchase_order_id": purchase_order_id,
        "total_cost": total_cost,
        "expected_arrival": expected_arrival,
        "status": status,
    }
