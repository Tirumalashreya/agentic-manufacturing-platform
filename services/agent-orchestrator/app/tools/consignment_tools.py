# app/tools/consignment_tools.py
import datetime
import openpyxl

from app.tools.audit_tools import record_audit_event
from app.tools.billing_tools import create_invoice
from app.tools.production_tools import reserve_stock

MOCK_ERP_PATH = "mock_erp.xlsx"

CONSIGNMENT_HEADERS = [
    "consignment_id",
    "customer_id",
    "customer_name",
    "product_id",
    "product_name",
    "unit_price",
    "quantity_sent",
    "quantity_sold",
    "quantity_returned",
    "quantity_remaining",
    "reservation_id",
    "notes",
    "status",
    "created_at",
    "closed_at",
]


def _consignments_sheet(wb):
    if "ConsignmentStock" not in wb.sheetnames:
        sheet = wb.create_sheet("ConsignmentStock")
        sheet.append(CONSIGNMENT_HEADERS)
        return sheet
    return wb["ConsignmentStock"]


def _next_consignment_id(sheet) -> str:
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith("CNG"):
            try:
                max_num = max(max_num, int(str(value)[3:]))
            except ValueError:
                pass
    return f"CNG{max_num + 1:04d}"


def find_consignment(consignment_id: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    if "ConsignmentStock" not in wb.sheetnames:
        return None
    sheet = wb["ConsignmentStock"]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == consignment_id:
            return dict(zip(CONSIGNMENT_HEADERS, row))
    return None


def _set_fields(consignment_id: str, **fields) -> None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = wb["ConsignmentStock"]
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value == consignment_id:
            for name, value in fields.items():
                sheet.cell(row=row_idx, column=CONSIGNMENT_HEADERS.index(name) + 1, value=value)
            wb.save(MOCK_ERP_PATH)
            return


# The critical rule this whole entity exists to enforce: transferring
# stock to a consignee is NOT a sale — reserve_stock() is reused for its
# real mechanic (atomically decrement real Inventory, fail clean if there
# isn't enough), but deliberately with no create_invoice() call anywhere
# near it. Nothing is billed until the consignee actually reports a sale.
def create_consignment(
    customer_id: str | None,
    customer_name: str | None,
    product_id: str | None,
    product_name: str | None,
    quantity_sent: int | None,
    unit_price: float | None,
    notes: str,
) -> dict:
    if product_id is None or unit_price is None:
        return {"created": False, "reason": "Cannot create a consignment without a valid product and price."}
    if quantity_sent is None or quantity_sent <= 0:
        return {"created": False, "reason": "Cannot create a consignment for zero or negative quantity."}

    reservation = reserve_stock(
        sales_order_id=None,
        product_id=product_id,
        product_name=product_name,
        quantity=quantity_sent,
    )
    if not reservation["reserved"]:
        return {"created": False, "reason": f"Stock transfer failed: {reservation['reason']}"}

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _consignments_sheet(wb)
    consignment_id = _next_consignment_id(sheet)
    created_at = datetime.datetime.now().isoformat()

    sheet.append([
        consignment_id, customer_id, customer_name, product_id, product_name,
        unit_price, quantity_sent, 0, 0, quantity_sent,
        reservation["reservation_id"], notes, "ACTIVE", created_at, None,
    ])
    wb.save(MOCK_ERP_PATH)
    record_audit_event("Consignment", consignment_id, "CONSIGNMENT_CREATED", {"quantity_sent": quantity_sent})

    return {
        "created": True,
        "consignment_id": consignment_id,
        "quantity_sent": quantity_sent,
        "status": "ACTIVE",
    }


# Invoices ONLY the quantity actually sold, not the full consignment —
# reuses create_invoice() exactly as the roadmap specified. Cumulative
# sold+returned can never exceed what was actually sent, same "can't
# exceed what's real" guard as post_sales_tools' return-quantity check.
def report_consignment_sale(consignment_id: str, quantity_sold: int) -> dict:
    if quantity_sold is None or quantity_sold <= 0:
        return {"reported": False, "reason": "Cannot report a sale for zero or negative quantity."}

    consignment = find_consignment(consignment_id)
    if consignment is None:
        return {"reported": False, "reason": f"No consignment found with id {consignment_id}"}
    if consignment["status"] != "ACTIVE":
        return {"reported": False, "reason": f"{consignment_id} is not ACTIVE (status: {consignment['status']})."}

    already_settled = consignment["quantity_sold"] + consignment["quantity_returned"]
    if already_settled + quantity_sold > consignment["quantity_sent"]:
        return {
            "reported": False,
            "reason": (
                f"Reported sale of {quantity_sold} plus {already_settled} already settled "
                f"would exceed the {consignment['quantity_sent']} units sent."
            ),
        }

    invoice = create_invoice(
        sales_order_id=None,
        delivery_id=None,
        customer_id=consignment["customer_id"],
        customer_name=consignment["customer_name"],
        quantity_billed=quantity_sold,
        unit_price=consignment["unit_price"],
        payment_terms_days=30,
        notes=f"Consignment {consignment_id} sell-through report: {quantity_sold} units sold.",
    )

    new_sold = consignment["quantity_sold"] + quantity_sold
    new_remaining = consignment["quantity_sent"] - new_sold - consignment["quantity_returned"]
    _set_fields(consignment_id, quantity_sold=new_sold, quantity_remaining=new_remaining)
    record_audit_event("Consignment", consignment_id, "CONSIGNMENT_SALE_REPORTED", {
        "quantity_sold": quantity_sold, "invoice_id": invoice["invoice_id"] if invoice["created"] else None,
    })

    return {
        "reported": True,
        "consignment_id": consignment_id,
        "invoice_id": invoice["invoice_id"] if invoice["created"] else None,
        "quantity_remaining": new_remaining,
    }


# Unsold stock physically comes back — credit real Inventory again, same
# "credit stock back on a real physical event" pattern as
# resolve_return_inspection()'s APPROVED branch.
def return_unsold_consignment_stock(consignment_id: str, quantity_returned: int) -> dict:
    if quantity_returned is None or quantity_returned <= 0:
        return {"returned": False, "reason": "Cannot return zero or negative quantity."}

    consignment = find_consignment(consignment_id)
    if consignment is None:
        return {"returned": False, "reason": f"No consignment found with id {consignment_id}"}
    if consignment["status"] != "ACTIVE":
        return {"returned": False, "reason": f"{consignment_id} is not ACTIVE (status: {consignment['status']})."}

    already_settled = consignment["quantity_sold"] + consignment["quantity_returned"]
    if already_settled + quantity_returned > consignment["quantity_sent"]:
        return {
            "returned": False,
            "reason": (
                f"Returning {quantity_returned} plus {already_settled} already settled "
                f"would exceed the {consignment['quantity_sent']} units sent."
            ),
        }

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    inventory = wb["Inventory"]
    for inv_row in inventory.iter_rows(min_row=2):
        pid_cell, _plant_cell, qty_cell = inv_row
        if pid_cell.value == consignment["product_id"]:
            qty_cell.value = (qty_cell.value or 0) + quantity_returned
            break
    wb.save(MOCK_ERP_PATH)

    new_returned = consignment["quantity_returned"] + quantity_returned
    new_remaining = consignment["quantity_sent"] - consignment["quantity_sold"] - new_returned
    _set_fields(consignment_id, quantity_returned=new_returned, quantity_remaining=new_remaining)
    record_audit_event("Consignment", consignment_id, "CONSIGNMENT_STOCK_RETURNED", {"quantity_returned": quantity_returned})

    return {"returned": True, "consignment_id": consignment_id, "quantity_remaining": new_remaining}


# Only closable once every unit sent has been accounted for — sold,
# returned, or a mix of both. Terminal.
def close_consignment(consignment_id: str) -> dict:
    consignment = find_consignment(consignment_id)
    if consignment is None:
        return {"closed": False, "reason": f"No consignment found with id {consignment_id}"}
    if consignment["status"] != "ACTIVE":
        return {"closed": False, "reason": f"{consignment_id} is not ACTIVE (status: {consignment['status']})."}
    if consignment["quantity_remaining"] > 0:
        return {
            "closed": False,
            "reason": f"Cannot close {consignment_id} — {consignment['quantity_remaining']} units still unaccounted for.",
        }

    _set_fields(consignment_id, status="CLOSED", closed_at=datetime.datetime.now().isoformat())
    record_audit_event("Consignment", consignment_id, "CONSIGNMENT_CLOSED")
    return {"closed": True, "consignment_id": consignment_id, "status": "CLOSED"}
