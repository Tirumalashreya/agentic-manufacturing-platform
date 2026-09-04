# app/tools/approval_tools.py
import datetime
import openpyxl

MOCK_ERP_PATH = "mock_erp.xlsx"

APPROVAL_HEADERS = [
    "approval_id",
    "quotation_id",
    "original_price",
    "requested_price",
    "requested_discount_percent",
    "reason",
    "status",
    "requested_at",
    "resolved_at",
    "resolved_by",
]

_STATUS_COL = APPROVAL_HEADERS.index("status") + 1
_RESOLVED_AT_COL = APPROVAL_HEADERS.index("resolved_at") + 1
_RESOLVED_BY_COL = APPROVAL_HEADERS.index("resolved_by") + 1


def _approvals_sheet(wb):
    if "ApprovalRequests" not in wb.sheetnames:
        sheet = wb.create_sheet("ApprovalRequests")
        sheet.append(APPROVAL_HEADERS)
        return sheet
    return wb["ApprovalRequests"]


def _next_approval_id(sheet) -> str:
    # Based on the max existing numeric suffix, not row count — robust to
    # any row ever being deleted.
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[0]
        if value and str(value).startswith("A"):
            try:
                max_num = max(max_num, int(str(value)[1:]))
            except ValueError:
                pass
    return f"A{max_num + 1:04d}"


def _find_approval_row(sheet, approval_id: str):
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value == approval_id:
            return row_idx, row
    return None, None


# A discount that failed the pricing policy check doesn't just get
# dropped — it becomes a real pending record a human can act on later.
def create_approval_request(
    quotation_id: str,
    original_price: float,
    requested_price: float,
    requested_discount_percent: float,
    reason: str,
) -> dict:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _approvals_sheet(wb)

    approval_id = _next_approval_id(sheet)
    requested_at = datetime.datetime.now().isoformat()
    status = "PENDING"

    sheet.append([
        approval_id,
        quotation_id,
        original_price,
        requested_price,
        requested_discount_percent,
        reason,
        status,
        requested_at,
        None,
        None,
    ])
    wb.save(MOCK_ERP_PATH)

    return {
        "created": True,
        "approval_id": approval_id,
        "status": status,
    }


def get_approval_request(approval_id: str) -> dict | None:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _approvals_sheet(wb)
    _, row = _find_approval_row(sheet, approval_id)
    if row is None:
        return None
    return dict(zip(APPROVAL_HEADERS, (cell.value for cell in row)))


# The human-in-the-loop step. Not called by an agent — this is the
# deterministic action a person takes, separate from any LLM reasoning.
def resolve_approval_request(
    approval_id: str,
    decision: str,
    resolved_by: str = "SYSTEM",
) -> dict:
    if decision not in ("APPROVED", "REJECTED"):
        return {"resolved": False, "reason": f"Invalid decision: {decision}"}

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = _approvals_sheet(wb)
    row_idx, row = _find_approval_row(sheet, approval_id)
    if row is None:
        return {"resolved": False, "reason": f"No approval request found with id {approval_id}"}

    current_status = row[_STATUS_COL - 1].value
    if current_status != "PENDING":
        return {
            "resolved": False,
            "reason": f"Approval request {approval_id} is already {current_status}, cannot resolve again.",
        }

    resolved_at = datetime.datetime.now().isoformat()
    sheet.cell(row=row_idx, column=_STATUS_COL, value=decision)
    sheet.cell(row=row_idx, column=_RESOLVED_AT_COL, value=resolved_at)
    sheet.cell(row=row_idx, column=_RESOLVED_BY_COL, value=resolved_by)
    wb.save(MOCK_ERP_PATH)

    return {
        "resolved": True,
        "approval_id": approval_id,
        "status": decision,
        "quotation_id": row[1].value,
        "requested_price": row[3].value,
    }
