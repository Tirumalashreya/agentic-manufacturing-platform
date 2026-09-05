# app/tests/test_post_sales.py
import openpyxl
import pytest

from app.agents.post_sales_agent import post_sales_agent
from app.tools.post_sales_tools import (
    MOCK_ERP_PATH,
    find_sales_order,
    get_post_sales_action,
    process_post_sales_action,
    resolve_return_inspection,
)
from app.tools.sales_tools import create_sales_order


@pytest.mark.asyncio
async def test_post_sales_agent_extracts_return():
    result = await post_sales_agent(
        process="PRODUCT_RETURN",
        user_message="ABC wants to return 5 laptops from order SO123.",
    )

    assert result.action == "RETURN"
    assert result.referenced_order_id == "SO123"


@pytest.mark.asyncio
async def test_post_sales_agent_extracts_credit_memo_no_order():
    result = await post_sales_agent(
        process="CREDIT_MEMO",
        user_message="We'd like a credit memo issued to our account.",
    )

    assert result.action == "CREDIT_MEMO"
    assert result.referenced_order_id is None


def _inventory_qty(product_id: str, plant: str) -> int:
    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = wb["Inventory"]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == product_id and row[1] == plant:
            return row[2]
    raise AssertionError(f"No inventory row for {product_id}/{plant}")


def test_find_sales_order_real_order():
    created = create_sales_order(
        customer_id="C001",
        customer_name="ABC Corp",
        customer_active=True,
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=3,
        unit_price=800.0,
        requested_deadline=None,
    )

    found = find_sales_order(created["sales_order_id"])

    assert found is not None
    assert found["customer_id"] == "C001"
    assert found["product_id"] == "P011"


def test_find_sales_order_unknown_id():
    assert find_sales_order("SO_DOES_NOT_EXIST") is None


def test_process_return_puts_stock_back():
    # Reference a real Sales Order — the previous version of this test
    # used a fake "SO123" that never existed, which the new reference
    # validation now correctly rejects.
    order = create_sales_order(
        customer_id="C001",
        customer_name="ABC Corp",
        customer_active=True,
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=3,
        unit_price=800.0,
        requested_deadline=None,
    )
    before = _inventory_qty("P011", "PLANT01")

    result = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=3,
        unit_price=800.0,
        reason="Customer changed their mind.",
    )

    assert result["created"] is True
    assert result["post_sales_id"].startswith("RET")
    assert result["amount"] == 2400.0
    assert result["status"] == "PENDING_INSPECTION"

    # Stock does NOT move yet — only after inspection approves it.
    after_request = _inventory_qty("P011", "PLANT01")
    assert after_request == before

    resolved = resolve_return_inspection(result["post_sales_id"], "APPROVED", inspected_by="qc_team")
    assert resolved["resolved"] is True
    assert resolved["status"] == "APPROVED"

    after_approval = _inventory_qty("P011", "PLANT01")
    assert after_approval - before == 3


def test_resolve_return_inspection_rejected_does_not_move_stock():
    order = create_sales_order(
        customer_id="C001",
        customer_name="ABC Corp",
        customer_active=True,
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=3,
        unit_price=800.0,
        requested_deadline=None,
    )
    before = _inventory_qty("P011", "PLANT01")

    created = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=2,
        unit_price=800.0,
        reason="Damaged on return.",
    )

    resolved = resolve_return_inspection(created["post_sales_id"], "REJECTED", inspected_by="qc_team")

    assert resolved["resolved"] is True
    assert resolved["status"] == "REJECTED"

    after = _inventory_qty("P011", "PLANT01")
    assert after == before

    fetched = get_post_sales_action(created["post_sales_id"])
    assert fetched["status"] == "REJECTED"
    assert fetched["inspected_by"] == "qc_team"


def test_resolve_return_inspection_cannot_double_resolve():
    order = create_sales_order(
        customer_id="C001",
        customer_name="ABC Corp",
        customer_active=True,
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=3,
        unit_price=800.0,
        requested_deadline=None,
    )
    created = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=1,
        unit_price=800.0,
        reason="Wrong color.",
    )

    resolve_return_inspection(created["post_sales_id"], "APPROVED")
    second = resolve_return_inspection(created["post_sales_id"], "REJECTED")

    assert second["resolved"] is False


def test_resolve_return_inspection_unknown_id():
    result = resolve_return_inspection("RET9999", "APPROVED")

    assert result["resolved"] is False


def test_resolve_return_inspection_invalid_decision():
    result = resolve_return_inspection("RET0001", "MAYBE")

    assert result["resolved"] is False


def _backdate_order(sales_order_id: str, days_ago: int) -> None:
    import datetime

    wb = openpyxl.load_workbook(MOCK_ERP_PATH)
    sheet = wb["SalesOrders"]
    created_at_col = 11  # 1-indexed: sales_order_id..status(10), created_at(11)
    backdated = (datetime.date.today() - datetime.timedelta(days=days_ago)).isoformat()
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row[0].value == sales_order_id:
            sheet.cell(row=row_idx, column=created_at_col, value=backdated)
            break
    wb.save(MOCK_ERP_PATH)


def test_process_post_sales_action_rejects_return_outside_window():
    order = create_sales_order(
        customer_id="C001",
        customer_name="ABC Corp",
        customer_active=True,
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=3,
        unit_price=800.0,
        requested_deadline=None,
    )
    _backdate_order(order["sales_order_id"], days_ago=40)  # beyond the 30-day window

    result = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=1,
        unit_price=800.0,
        reason="Changed my mind, late.",
    )

    assert result["created"] is False
    assert "return window" in result["reason"]


def test_process_post_sales_action_allows_return_within_window():
    order = create_sales_order(
        customer_id="C001",
        customer_name="ABC Corp",
        customer_active=True,
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=3,
        unit_price=800.0,
        requested_deadline=None,
    )
    _backdate_order(order["sales_order_id"], days_ago=10)  # within the 30-day window

    result = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=1,
        unit_price=800.0,
        reason="Changed my mind, still in time.",
    )

    assert result["created"] is True


def test_process_post_sales_action_rejects_unknown_reference():
    result = process_post_sales_action(
        action="RETURN",
        referenced_order_id="SO_DOES_NOT_EXIST",
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=1,
        unit_price=800.0,
        reason="Customer changed their mind.",
    )

    assert result["created"] is False
    assert "does not exist" in result["reason"]


def test_process_post_sales_action_rejects_customer_mismatch():
    order = create_sales_order(
        customer_id="C001",
        customer_name="ABC Corp",
        customer_active=True,
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=5,
        unit_price=800.0,
        requested_deadline=None,
    )

    # Order belongs to C001 — a different customer has no standing to
    # return against it, even with a matching product/quantity.
    result = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C002",
        customer_name="Zenith Enterprises",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=1,
        unit_price=800.0,
        reason="Not my order but trying anyway.",
    )

    assert result["created"] is False
    assert "does not belong to customer" in result["reason"]


def test_process_post_sales_action_rejects_product_not_on_order():
    order = create_sales_order(
        customer_id="C001",
        customer_name="ABC Corp",
        customer_active=True,
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=5,
        unit_price=800.0,
        requested_deadline=None,
    )

    # Order was for Wireless Mouse (P011); trying to return a Laptop
    # (P001) against it should be rejected.
    result = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P001",
        product_name="Laptop - Standard",
        quantity=1,
        unit_price=50000.0,
        reason="Wrong item.",
    )

    assert result["created"] is False
    assert "was not part of order" in result["reason"]


def test_process_post_sales_action_rejects_quantity_exceeding_order():
    order = create_sales_order(
        customer_id="C001",
        customer_name="ABC Corp",
        customer_active=True,
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=5,
        unit_price=800.0,
        requested_deadline=None,
    )

    # Only 5 were ordered — trying to return 10 should be rejected.
    result = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=10,
        unit_price=800.0,
        reason="Too many, wants to return extra.",
    )

    assert result["created"] is False
    assert "exceeds the" in result["reason"]


def test_process_post_sales_action_allows_quantity_within_order():
    order = create_sales_order(
        customer_id="C001",
        customer_name="ABC Corp",
        customer_active=True,
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=5,
        unit_price=800.0,
        requested_deadline=None,
    )

    result = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=2,
        unit_price=800.0,
        reason="Two units didn't fit the need.",
    )

    assert result["created"] is True


def test_process_post_sales_action_rejects_cumulative_over_return():
    order = create_sales_order(
        customer_id="C001",
        customer_name="ABC Corp",
        customer_active=True,
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=5,
        unit_price=800.0,
        requested_deadline=None,
    )

    # First return of 3 is individually within the 5-unit order — allowed.
    first = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=3,
        unit_price=800.0,
        reason="First batch didn't work.",
    )
    assert first["created"] is True

    # Second return of 3 is ALSO individually within the 5-unit order, but
    # 3 (already returned) + 3 (this request) = 6 > 5 ordered. Must reject.
    second = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=3,
        unit_price=800.0,
        reason="Second batch also didn't work.",
    )
    assert second["created"] is False
    assert "already returned" in second["reason"]

    # A second return that stays within the remaining balance (2 left)
    # should still succeed.
    third = process_post_sales_action(
        action="RETURN",
        referenced_order_id=order["sales_order_id"],
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=2,
        unit_price=800.0,
        reason="Remaining units also didn't work.",
    )
    assert third["created"] is True


def test_process_credit_memo_does_not_touch_inventory():
    before = _inventory_qty("P011", "PLANT01")

    result = process_post_sales_action(
        action="CREDIT_MEMO",
        referenced_order_id=None,
        customer_id="C001",
        customer_name="ABC Corp",
        product_id="P011",
        product_name="Wireless Mouse",
        quantity=2,
        unit_price=800.0,
        reason="Goodwill credit.",
    )

    assert result["created"] is True
    assert result["post_sales_id"].startswith("CM")

    after = _inventory_qty("P011", "PLANT01")
    assert after == before


def test_process_post_sales_action_rejects_missing_product():
    result = process_post_sales_action(
        action="RETURN",
        referenced_order_id=None,
        customer_id="C001",
        customer_name="ABC Corp",
        product_id=None,
        product_name=None,
        quantity=1,
        unit_price=None,
        reason="",
    )

    assert result["created"] is False
